"""Hidden draft identity inside an exported DOCX, and the verified read-back.

"Word'de düzelttim, geri yükle" (W22). A lawyer exports a draft to Word,
edits it there, and uploads the edited file back. Competitors ship a Word
add-in; ColleX does a VERIFIED round trip instead: the returned document is
matched to the draft it came from, compared paragraph by paragraph, and saved
only through the same revise path (``reviseDraft``) as every console edit —
so the quote-integrity gate, the KAYNAKSIZ discipline and the locked sections
apply to a Word edit exactly as they apply to any other.

WHAT THE EXPORT EMBEDS (every mode — TASLAK and NİHAİ alike). Nothing here is
visible text: the printed document and every paragraph's ``.text`` are
unchanged.

1. A **custom XML part** (``/customXml/itemN.xml`` + its ``itemProps``),
   related from the main document part — the OOXML data store Word keeps on
   save. It carries ``collex.draft-identity/v1``: the draft id, the version,
   the template, the export mode, and per rendered draft paragraph its id,
   section, role, ``supported`` flag, the sha256 of its canonical text and,
   for every evidence it cites, the ``quoteSha256``. A ``digest`` attribute
   (sha256 over a fixed serialization of all of that) makes the part
   **tamper-evident**: an edited attribute no longer matches its digest.
   It is NOT a signature — anyone can recompute a sha256 — and it does not
   need to be one: an import never grants anything the console's own
   ``PUT /v1/drafts/{id}`` would not; every gate runs again on the server.
2. A **hidden bookmark** (``_cxp<n>``; a leading underscore is Word's hidden
   bookmark convention, and the 40-character name limit is respected) around
   the text lines of every rendered draft paragraph. Bookmarks travel with
   the text when the lawyer edits it, so a changed paragraph is still found
   by its id; where a bookmark was lost the importer aligns by text.

WHAT THE READ-BACK RETURNS (``read_draft_docx``, CLI below). The identity and
whether its digest holds; every body block (paragraph / heading / table) in
document order with its style name, its bookmark numbers and its ACCEPTED
text — tracked insertions in, tracked deletions out (the same rule
``intake/extract.py`` applies to an uploaded DOCX) — and a count of the
tracked changes still pending, so the lawyer is TOLD the text was read with
them accepted. The alignment and the diff are computed by the control-plane
(``control-plane/src/drafting/docxImport.ts``), next to the stored draft.

CLI (used by ``POST /v1/drafts/{id}/import-docx``; JSON on stdout)::

    python -X utf8 -m export.draft_identity --read <file.docx> --name <original name>

Exit 0 with the read-back; exit 1 with ``{"error": {"kind", "message"}}``
(``UNSUPPORTED_TYPE`` / ``INVALID_REQUEST`` / ``EXTRACTION_FAILED``) — the
upload quarantine (``intake/quarantine.py::verify_upload``: magic bytes,
ZIP-bomb caps, the 25 MiB cap) runs BEFORE any XML is parsed.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable

from export.draft import (
    EK_DOGRULAMA_SECTION_ID,
    Draft,
    DraftParagraph,
    DraftSection,
    ExportMode,
    canonical_quote_text,
    sha256_utf8,
)

IDENTITY_SCHEMA = "collex.draft-identity/v1"
IDENTITY_NS = "urn:collex:draft-identity:1"
READBACK_SCHEMA = "collex.draft-docx-readback/v1"

#: Hidden bookmark name prefix: ``_cxp`` + the paragraph's 1-based render
#: number. Word hides bookmarks whose name starts with "_" and truncates names
#: past 40 characters; ``_cxp`` + a number stays far below that.
BOOKMARK_PREFIX = "_cxp"

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_W = "{" + _W_NS + "}"
_DS_NS = "http://schemas.openxmlformats.org/officeDocument/2006/customXml"
_CUSTOM_XML_PROPS_CT = "application/vnd.openxmlformats-officedocument.customXmlProperties+xml"
_RT_CUSTOM_XML = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml"
_RT_CUSTOM_XML_PROPS = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXmlProps"
)
_ITEM_PARTNAME = re.compile(r"\A/customXml/item\d+\.xml\Z", re.IGNORECASE)
_BOOKMARK_NAME = re.compile(r"\A_cxp(\d{1,6})\Z")
#: Namespace for the deterministic datastore item GUID (uuid5).
_ITEM_UUID_NS = uuid.UUID("8f2f7a0e-0b8e-4c55-9d7c-3f1c2b7d9a41")

#: Identity statuses the read-back reports (the TS route maps each to its own
#: Turkish refusal; only OK proceeds).
STATUS_OK = "OK"
STATUS_MISSING = "MISSING"
STATUS_DIGEST_MISMATCH = "DIGEST_MISMATCH"
STATUS_MALFORMED = "MALFORMED"
STATUS_AMBIGUOUS = "AMBIGUOUS"


# --------------------------------------------------------------------------
# Identity model
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class IdentityQuote:
    evidence_id: str
    quote_sha256: str


@dataclass(frozen=True)
class IdentityParagraph:
    n: int
    paragraph_id: str
    section_id: str
    role: str
    supported: bool
    lines: int
    text_sha256: str
    quotes: tuple[IdentityQuote, ...]


@dataclass(frozen=True)
class DraftIdentity:
    draft_id: str
    version: int
    template: str
    annex: str
    marks: str
    paragraphs: tuple[IdentityParagraph, ...]
    #: sha256 (canonical text) of every NON-bookmarked, non-empty paragraph the
    #: exporter wrote: banners, headings, citation and note lines, the
    #: appendix. The importer skips a paragraph whose fingerprint is here, so
    #: machine apparatus is never mistaken for a lawyer's new paragraph — and
    #: an apparatus line that no longer matches is REPORTED, never applied.
    apparatus: tuple[str, ...] = ()

    def digest(self) -> str:
        """sha256 over a FIXED serialization — the tamper-evidence value."""
        lines = [
            IDENTITY_SCHEMA,
            f"draftId={self.draft_id}",
            f"version={self.version}",
            f"template={self.template}",
            f"annex={self.annex}",
            f"marks={self.marks}",
        ]
        for p in self.paragraphs:
            quotes = ";".join(f"{q.evidence_id}:{q.quote_sha256}" for q in p.quotes)
            lines.append(
                "|".join(
                    [
                        "p",
                        str(p.n),
                        p.paragraph_id,
                        p.section_id,
                        p.role,
                        "1" if p.supported else "0",
                        str(p.lines),
                        p.text_sha256,
                        quotes,
                    ]
                )
            )
        lines.append("apparatus=" + ",".join(self.apparatus))
        return sha256_utf8("\n".join(lines))

    def to_json_dict(self) -> dict[str, Any]:
        return {
            "schema": IDENTITY_SCHEMA,
            "draftId": self.draft_id,
            "version": self.version,
            "template": self.template,
            "exportMode": {"annex": self.annex, "marks": self.marks},
            "paragraphs": [
                {
                    "n": p.n,
                    "id": p.paragraph_id,
                    "sectionId": p.section_id,
                    "role": p.role,
                    "supported": p.supported,
                    "lines": p.lines,
                    "textSha256": p.text_sha256,
                    "quotes": [
                        {"evidenceId": q.evidence_id, "quoteSha256": q.quote_sha256}
                        for q in p.quotes
                    ],
                }
                for p in self.paragraphs
            ],
            "apparatusSha256": list(self.apparatus),
        }


def paragraph_text_sha256(text: str) -> str:
    """sha256 of the CANONICAL paragraph text (the quote-integrity form).

    The control-plane recomputes the same value with ``canonicalQuoteText``
    (the two canonicalizations are pinned to each other), so an identity can
    be checked against the stored version without sharing any other code.
    """
    return sha256_utf8(canonical_quote_text(text))


def build_identity(
    draft: Draft,
    mode: ExportMode,
    rendered: Iterable[tuple[DraftSection, DraftParagraph, int]],
    apparatus: Iterable[str] = (),
) -> DraftIdentity:
    """The identity of the paragraphs that were actually rendered, in order."""
    by_id = draft.evidence_by_id()
    paragraphs: list[IdentityParagraph] = []
    for n, (section, paragraph, lines) in enumerate(rendered, start=1):
        quotes = tuple(
            IdentityQuote(evidence_id, by_id[evidence_id].quote_sha256)
            for evidence_id in paragraph.evidence_ids
            if evidence_id in by_id
        )
        paragraphs.append(
            IdentityParagraph(
                n=n,
                paragraph_id=paragraph.paragraph_id,
                section_id=section.section_id,
                role=paragraph.role,
                supported=paragraph.supported,
                lines=lines,
                text_sha256=paragraph_text_sha256(paragraph.text),
                quotes=quotes,
            )
        )
    return DraftIdentity(
        draft_id=draft.draft_id,
        version=draft.version,
        template=draft.template,
        annex=mode.annex,
        marks=mode.marks,
        paragraphs=tuple(paragraphs),
        apparatus=tuple(sorted(set(apparatus))),
    )


def identity_xml(identity: DraftIdentity) -> bytes:
    """Serialize the identity as the custom XML part's payload."""
    from lxml import etree

    root = etree.Element(f"{{{IDENTITY_NS}}}draftIdentity", nsmap={"cx": IDENTITY_NS})
    root.set("schema", IDENTITY_SCHEMA)
    root.set("draftId", identity.draft_id)
    root.set("version", str(identity.version))
    root.set("template", identity.template)
    root.set("annex", identity.annex)
    root.set("marks", identity.marks)
    root.set("digest", identity.digest())
    for p in identity.paragraphs:
        node = etree.SubElement(root, f"{{{IDENTITY_NS}}}p")
        node.set("n", str(p.n))
        node.set("id", p.paragraph_id)
        node.set("section", p.section_id)
        node.set("role", p.role)
        node.set("supported", "true" if p.supported else "false")
        node.set("lines", str(p.lines))
        node.set("textSha256", p.text_sha256)
        for q in p.quotes:
            quote = etree.SubElement(node, f"{{{IDENTITY_NS}}}quote")
            quote.set("evidenceId", q.evidence_id)
            quote.set("quoteSha256", q.quote_sha256)
    for digest in identity.apparatus:
        etree.SubElement(root, f"{{{IDENTITY_NS}}}apparatus").set("textSha256", digest)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _item_props_xml(item_id: str) -> bytes:
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="no"?>\n'
        f'<ds:datastoreItem ds:itemID="{{{item_id}}}" xmlns:ds="{_DS_NS}">'
        f'<ds:schemaRefs><ds:schemaRef ds:uri="{IDENTITY_NS}"/></ds:schemaRefs>'
        "</ds:datastoreItem>"
    ).encode("utf-8")


# --------------------------------------------------------------------------
# Embedding (called by export/petition.py::build_petition_document)
# --------------------------------------------------------------------------


def _add_bookmark(first: Any, last: Any, n: int) -> None:
    """Hidden bookmark ``_cxp<n>`` around the paragraph's rendered lines."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(n))
    start.set(qn("w:name"), f"{BOOKMARK_PREFIX}{n}")
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(n))
    p_first = first._p
    ppr = p_first.pPr
    if ppr is not None:
        ppr.addnext(start)
    else:
        p_first.insert(0, start)
    last._p.append(end)


def attach_draft_identity(
    document: Any,
    draft: Draft,
    mode: ExportMode,
    rendered: list[tuple[DraftSection, DraftParagraph, list[Any]]],
) -> DraftIdentity:
    """Bookmark every rendered paragraph and add the identity part.

    ``rendered`` lists, in document order, each draft paragraph that was
    written together with the python-docx paragraphs of its TEXT lines (never
    its citation or note lines — those are regenerated, not edited).
    """
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    bookmarked: set[int] = set()
    for n, (_section, _paragraph, lines) in enumerate(rendered, start=1):
        if lines:
            _add_bookmark(lines[0], lines[-1], n)
        bookmarked.update(id(line._p) for line in lines)
    apparatus = {
        paragraph_text_sha256(paragraph.text)
        for paragraph in document.paragraphs
        if id(paragraph._p) not in bookmarked and canonical_quote_text(paragraph.text) != ""
    }
    identity = build_identity(
        draft,
        mode,
        ((section, paragraph, len(lines)) for section, paragraph, lines in rendered),
        apparatus,
    )
    package = document.part.package
    item_name = package.next_partname("/customXml/item%d.xml")
    number = re.search(r"(\d+)\.xml\Z", str(item_name))
    suffix = number.group(1) if number is not None else "1"
    item = Part(PackURI(str(item_name)), "application/xml", identity_xml(identity), package)
    item_id = str(uuid.uuid5(_ITEM_UUID_NS, f"{draft.draft_id}/{draft.version}")).upper()
    props = Part(
        PackURI(f"/customXml/itemProps{suffix}.xml"),
        _CUSTOM_XML_PROPS_CT,
        _item_props_xml(item_id),
        package,
    )
    document.part.relate_to(item, _RT_CUSTOM_XML)
    item.relate_to(props, _RT_CUSTOM_XML_PROPS)
    return identity


# --------------------------------------------------------------------------
# Read-back
# --------------------------------------------------------------------------


def _safe_root(blob: bytes) -> Any:
    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True, huge_tree=False)
    return etree.fromstring(blob, parser)


def _parse_identity(root: Any) -> tuple[DraftIdentity | None, str]:
    if root.get("schema") != IDENTITY_SCHEMA:
        return None, STATUS_MALFORMED
    try:
        version = int(root.get("version", ""))
    except ValueError:
        return None, STATUS_MALFORMED
    paragraphs: list[IdentityParagraph] = []
    for node in root.findall(f"{{{IDENTITY_NS}}}p"):
        try:
            n = int(node.get("n", ""))
            lines = int(node.get("lines", ""))
        except ValueError:
            return None, STATUS_MALFORMED
        quotes = tuple(
            IdentityQuote(q.get("evidenceId", ""), q.get("quoteSha256", ""))
            for q in node.findall(f"{{{IDENTITY_NS}}}quote")
        )
        paragraphs.append(
            IdentityParagraph(
                n=n,
                paragraph_id=node.get("id", ""),
                section_id=node.get("section", ""),
                role=node.get("role", ""),
                supported=node.get("supported") == "true",
                lines=lines,
                text_sha256=node.get("textSha256", ""),
                quotes=quotes,
            )
        )
    identity = DraftIdentity(
        draft_id=root.get("draftId", ""),
        version=version,
        template=root.get("template", ""),
        annex=root.get("annex", ""),
        marks=root.get("marks", ""),
        paragraphs=tuple(paragraphs),
        apparatus=tuple(
            node.get("textSha256", "") for node in root.findall(f"{{{IDENTITY_NS}}}apparatus")
        ),
    )
    if identity.draft_id == "" or any(p.paragraph_id == "" for p in paragraphs):
        return None, STATUS_MALFORMED
    if root.get("digest") != identity.digest():
        # The identity is returned anyway (for the diagnostic line) but the
        # status refuses it: an edited identity is not an identity.
        return identity, STATUS_DIGEST_MISMATCH
    return identity, STATUS_OK


def read_identity(document: Any) -> tuple[DraftIdentity | None, str]:
    """Find the ColleX identity part, wherever Word renumbered it."""
    found: list[tuple[DraftIdentity | None, str]] = []
    for part in document.part.package.iter_parts():
        if not _ITEM_PARTNAME.match(str(part.partname)):
            continue
        try:
            root = _safe_root(part.blob)
        except Exception:  # noqa: BLE001 — a foreign, unparsable item is not ours
            continue
        if root.tag != f"{{{IDENTITY_NS}}}draftIdentity":
            continue
        found.append(_parse_identity(root))
    if not found:
        return None, STATUS_MISSING
    if len(found) > 1:
        return None, STATUS_AMBIGUOUS
    return found[0]


# Accepted-text reading is NOT reimplemented here: it is the one reader
# intake/extract.py applies to every uploaded DOCX (W22) — the text reads as if
# every tracked change had been ACCEPTED (insertions in, deletions and the
# source side of a move out). One rule, one source. Imported lazily so an
# EXPORT never loads the intake stack (pypdf, OCR probing) it does not use.
def accepted_paragraph_text(p_el: Any) -> str:
    from intake.extract import _docx_paragraph_text

    return _docx_paragraph_text(p_el)

_REVISION_TAGS = (_W + "ins", _W + "del", _W + "moveFrom", _W + "moveTo")
_FORMAT_REVISION_TAGS = (_W + "rPrChange", _W + "pPrChange")


def _paragraph_mark_deleted(p_el: Any) -> bool:
    """A tracked deletion of the paragraph MARK: accepted, it joins the next."""
    ppr = p_el.find(_W + "pPr")
    if ppr is None:
        return False
    rpr = ppr.find(_W + "rPr")
    return rpr is not None and rpr.find(_W + "del") is not None


def _body_blocks(container: Any) -> list[Any]:
    """Body-level w:p / w:tbl / bookmark markers in order (through w:sdt)."""
    blocks: list[Any] = []
    for child in container:
        tag = child.tag
        if tag in (_W + "p", _W + "tbl", _W + "bookmarkStart", _W + "bookmarkEnd"):
            blocks.append(child)
        elif tag == _W + "sdt":
            content = child.find(_W + "sdtContent")
            if content is not None:
                blocks.extend(_body_blocks(content))
        elif tag == _W + "customXml":
            blocks.extend(_body_blocks(child))
    return blocks


def _style_names(document: Any) -> dict[str, str]:
    names: dict[str, str] = {}
    for style in document.styles:
        style_id = getattr(style, "style_id", None)
        if style_id:
            names[style_id] = style.name or style_id
    return names


def _bookmark_number(name: str | None) -> int | None:
    match = _BOOKMARK_NAME.match(name or "")
    return int(match.group(1)) if match is not None else None


@dataclass
class _Block:
    kind: str
    style: str
    text: str
    marks: list[int]
    pending: bool

    def to_json_dict(self, index: int) -> dict[str, Any]:
        return {
            "index": index,
            "kind": self.kind,
            "style": self.style,
            "text": self.text,
            "marks": self.marks,
            "pendingChange": self.pending,
        }


def read_body_blocks(document: Any) -> tuple[list[_Block], dict[str, int]]:
    """Body blocks as they read with every tracked change accepted."""
    styles = _style_names(document)
    open_marks: dict[str, int] = {}  # bookmark w:id -> render number
    # A bookmark name is unique in a Word document; a SECOND start of the same
    # _cxp<n> (a copied paragraph, or a tool that duplicated the markup) is not
    # the paragraph's identity. It is ignored, so the copy aligns by its text
    # (W23: a copied list item was read as part of the item it was copied from).
    started: set[int] = set()
    blocks: list[_Block] = []
    counts = {"insertions": 0, "deletions": 0, "moves": 0, "formatting": 0}
    carry: _Block | None = None

    def count_revisions(el: Any) -> bool:
        pending = False
        for node in el.iter(*_REVISION_TAGS):
            pending = True
            if node.tag == _W + "ins":
                counts["insertions"] += 1
            elif node.tag == _W + "del":
                counts["deletions"] += 1
            else:
                counts["moves"] += 1
        for _ in el.iter(*_FORMAT_REVISION_TAGS):
            counts["formatting"] += 1
        return pending

    for element in _body_blocks(document.element.body):
        tag = element.tag
        if tag == _W + "bookmarkStart":
            number = _bookmark_number(element.get(_W + "name"))
            if number is not None and number not in started:
                started.add(number)
                open_marks[element.get(_W + "id", "")] = number
            continue
        if tag == _W + "bookmarkEnd":
            open_marks.pop(element.get(_W + "id", ""), None)
            continue
        if tag == _W + "tbl":
            pending = count_revisions(element)
            blocks.append(_Block("table", "", "", [], pending))
            continue
        # A paragraph: the marks open when it starts plus the ones it opens.
        marks = list(open_marks.values())
        closed: list[str] = []
        for node in element.iter(_W + "bookmarkStart", _W + "bookmarkEnd"):
            if node.tag == _W + "bookmarkStart":
                number = _bookmark_number(node.get(_W + "name"))
                if number is not None and number not in started:
                    started.add(number)
                    open_marks[node.get(_W + "id", "")] = number
                    if number not in marks:
                        marks.append(number)
            else:
                closed.append(node.get(_W + "id", ""))
        for mark_id in closed:
            open_marks.pop(mark_id, None)
        ppr = element.find(_W + "pPr")
        style_el = ppr.find(_W + "pStyle") if ppr is not None else None
        style_id = style_el.get(_W + "val") if style_el is not None else None
        style = styles.get(style_id or "", style_id or "Normal") if style_id else "Normal"
        kind = "heading" if style.startswith("Heading") or style == "Title" else "paragraph"
        pending = count_revisions(element)
        block = _Block(kind, style, accepted_paragraph_text(element), sorted(set(marks)), pending)
        if carry is not None:
            # The previous paragraph's mark was deleted (tracked): accepted,
            # the two read as ONE paragraph, in the earlier one's style.
            block = _Block(
                carry.kind,
                carry.style,
                carry.text + block.text,
                sorted(set(carry.marks) | set(block.marks)),
                True,
            )
            carry = None
        if _paragraph_mark_deleted(element):
            carry = block
            continue
        blocks.append(block)
    if carry is not None:
        blocks.append(carry)
    return blocks, counts


def read_draft_docx(source: str | Path | bytes) -> dict[str, Any]:
    """The read-back the control-plane aligns (``collex.draft-docx-readback/v1``)."""
    from docx import Document

    stream: Any = io.BytesIO(source) if isinstance(source, (bytes, bytearray)) else str(source)
    document = Document(stream)
    identity, status = read_identity(document)
    blocks, counts = read_body_blocks(document)
    pending_total = counts["insertions"] + counts["deletions"] + counts["moves"]
    return {
        "schema": READBACK_SCHEMA,
        "identityStatus": status,
        "identity": identity.to_json_dict() if identity is not None else None,
        "blocks": [block.to_json_dict(index) for index, block in enumerate(blocks)],
        "trackedChanges": {
            "pending": pending_total > 0,
            "insertions": counts["insertions"],
            "deletions": counts["deletions"],
            "moves": counts["moves"],
            "formatting": counts["formatting"],
        },
    }


# --------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------


def _error(kind: str, message: str) -> dict[str, Any]:
    return {"error": {"kind": kind, "message": message}}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m export.draft_identity")
    parser.add_argument("--read", required=True, help="edited DOCX to read back")
    parser.add_argument("--name", default="", help="the upload's original file name")
    args = parser.parse_args(argv)

    from intake.errors import IntakeError
    from intake.quarantine import verify_upload

    try:
        data = Path(args.read).read_bytes()
    except OSError:
        print(json.dumps(_error("INVALID_REQUEST", "yüklenen dosya okunamadı"), ensure_ascii=False))
        return 1
    try:
        verified = verify_upload(args.name or Path(args.read).name, data)
    except IntakeError as exc:
        print(json.dumps(exc.to_json_dict(), ensure_ascii=False))
        return 1
    if verified.kind != "docx":
        print(
            json.dumps(
                _error(
                    "UNSUPPORTED_TYPE",
                    "Geri yükleme yalnız Word belgesi (.docx) kabul eder;"
                    f" yüklenen dosya {verified.kind.upper()}.",
                ),
                ensure_ascii=False,
            )
        )
        return 1
    try:
        payload = read_draft_docx(data)
    except Exception as exc:  # noqa: BLE001 — a broken DOCX is a typed 422
        print(
            json.dumps(
                _error("EXTRACTION_FAILED", f"Word belgesi okunamadı ({type(exc).__name__})."),
                ensure_ascii=False,
            )
        )
        return 1
    payload["fileSha256"] = verified.sha256
    print(json.dumps(payload, ensure_ascii=False))
    return 0


if __name__ == "__main__":  # pragma: no cover
    sys.exit(main())
