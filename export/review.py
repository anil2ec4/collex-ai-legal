"""Sözleşme İnceleme Raporu → DOCX (``collex.contract-review/v1``) — W14 · B-24.

B-24's acceptance says the review "çıktı DOCX olarak iner". Phase A (L-EVID)
shipped the review ENGINE and the JSON endpoint and recorded, honestly, that
the DOCX half was not done (W14-L-EVID.md §13 item 1). This module is that
half, built to the pattern `export/audit.py` established for B-13: parse
strictly, build, write to a temporary name, re-open the artifact, self-check
what a reader will actually see, and only then publish.

THE THREE RULES THIS WRITER ENFORCES, and the self-check proves on the file:

1. **"yok" is a statement about the TEXT, never about the law.** The legend
   says so verbatim: a missing clause is not a defective contract, and only
   the lawyer decides whether it was needed. A checklist result is a search
   result.

2. **An unsourced observation never uses the word "risk".** The producer
   already strips every inflected form and prefixes the line with
   ``⚠ KAYNAKSIZ`` (clauseReview.ts ``stripRiskWords``); this writer refuses
   to publish a document in which a KAYNAKSIZ line contains the word anyway,
   so a hand-edited or replayed JSON cannot smuggle a risk assessment into a
   report that says it makes none.

3. **The report says it is rule-based.** No model ran, nothing was
   interpreted, and the banner at the top and bottom says exactly that. The
   danger with a document like this is not that it is wrong — it is that it
   looks like counsel's opinion.

2026-09-27 (contract-review defects on real contract shapes):

- every ``var`` / ``belirsiz`` row now carries the contract's OWN sentence
  (``findings[].matches[].excerpt``), printed under "Metinden alıntılar" and
  self-checked like the observations — without it the lawyer could not catch
  a wrong match;
- a clause is named by its display label (``clauses[].label``: "Özel
  Şartlar 1", "Madde 5", "Giriş") — "Madde #1" for the unnumbered opening
  block, and a bare "1" for a restarted numbering, are gone;
- the banner is printed at the top and the bottom and nowhere else, and the
  KAYNAKSIZ rule once: the report's ``notices`` repeat both sentences, so a
  notice already printed is not printed again (it used to appear three and
  two times);
- an observation the producer could not attach to a clause
  (``unattachedObservations``) is printed with its reason, never dropped.

Requires ``python-docx`` (the ``export`` extra), like the other DOCX writers.
"""

from __future__ import annotations

import json
import os
import re
import tempfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

from export import text as T
from export.errors import ExportError, ExportRefused

REVIEW_SCHEMA = "collex.contract-review/v1"
FORMAT_NAME = "inceleme-docx"

#: A4, same geometry contract as the petition and audit writers.
A4_WIDTH_CM = 21.0
A4_HEIGHT_CM = 29.7
MARGIN_CM = 2.0

TITLE = "SÖZLEŞME İNCELEME RAPORU"

#: Mirror of REVIEW_NOTICES[0] in control-plane/src/contracts/clauseReview.ts.
REVIEW_BANNER = (
    "Bu inceleme, sözleşme metnini sizin kontrol listenizle satır satır"
    " karşılaştırmaktan ibarettir. Hukukî değerlendirme yapmaz, yorum üretmez"
    " ve yapay zekâ kullanmaz."
)

#: The unsupported marker, identical to the drafting and review surfaces.
KAYNAKSIZ_PREFIX = "⚠ KAYNAKSIZ"

#: State legend — "yok" is about the TEXT, and the document says so.
STATE_LEGEND: tuple[tuple[str, str], ...] = (
    ("var", "Bu başlığı karşılayan bir madde bulundu (madde numarası yanında)."),
    (
        "yok",
        "Bu başlık, incelenen metinde BULUNAMADI. Bu, maddenin hukuken zorunlu"
        " olduğu ya da sözleşmenin sakat olduğu anlamına gelmez — değerlendirme"
        " avukatındır.",
    ),
    (
        "belirsiz",
        "Bir madde bu başlıkla ilişkili görünüyor ancak karşılayıp"
        " karşılamadığına karar verilemedi; maddeyi okuyun.",
    ),
)

#: The KAYNAKSIZ rule, printed where a reader can see it.
UNSOURCED_NOTE = (
    "Kaynağa bağlanamayan hiçbir gözlem 'risk' olarak adlandırılmaz;"
    f" {KAYNAKSIZ_PREFIX} etiketiyle ve gözlem olarak yazılır."
)

STYLE_TITLE = "CollexIncelemeBaslik"
STYLE_WARNING = "CollexIncelemeUyari"
STYLE_FIELD = "CollexIncelemeAlan"
STYLE_LEGEND = "CollexIncelemeSozluk"

_TABLE_HEADERS = ("Başlık", "Durum", "Madde", "Eşleşen ifade", "Not")

#: Every inflected form of "risk" (Turkish attaches suffixes to the stem).
_RISK_WORD = re.compile(r"\b[Rr][İiIı][Ss][Kk][^\W\d_]*\b", re.UNICODE)


class ReviewFormatError(ExportError):
    """The input JSON is not a valid ``collex.contract-review/v1`` report."""


@dataclass(frozen=True)
class ReviewExcerpt:
    """One contract sentence a finding rests on (``findings[].matches[]``)."""

    clause_label: str
    excerpt: str
    negated: bool
    weak: bool


@dataclass(frozen=True)
class ReviewFinding:
    label: str
    state_label: str
    clause_numbers: tuple[str, ...]
    matched_term: str
    note: str
    #: Display labels ("Özel Şartlar 1"); empty in reports written before them.
    clause_labels: tuple[str, ...] = ()
    #: Why a ``belirsiz`` row is ``belirsiz`` (Turkish sentence); "" otherwise.
    reason: str = ""
    matches: tuple[ReviewExcerpt, ...] = ()


@dataclass(frozen=True)
class ReviewObservation:
    clause_number: str
    text: str
    sourced: bool
    evidence_label: str
    #: Where the line is printed ("Madde 4", "Giriş", "Özel Şartlar 1").
    where: str = ""
    #: For an observation no clause could hold: why (printed with it).
    note: str = ""


@dataclass(frozen=True)
class ReviewReport:
    document_title: str
    checklist_title: str
    generated_at: str
    clause_count: int
    findings: tuple[ReviewFinding, ...]
    observations: tuple[ReviewObservation, ...]
    totals: dict[str, int]
    notices: tuple[str, ...]
    #: Observations whose ``clauseIndex`` named no clause (never dropped).
    unattached: tuple[ReviewObservation, ...] = ()


def clause_display_label(clause_number: str) -> str:
    """The printed name of a clause for a report that carries no ``label``.

    "#1" is the unnumbered opening block the producer keeps so nothing is
    dropped — "Madde #1" read like a clause the contract numbered.
    """
    if clause_number == "":
        return "Madde"
    if clause_number.startswith("#"):
        position = clause_number[1:]
        return "Giriş" if position == "1" else f"Numarasız bölüm {position}"
    return f"Madde {clause_number}"


def _require(payload: Any, key: str, *, where: str) -> Any:
    if key not in payload:
        raise ReviewFormatError(f"{where}: '{key}' alanı eksik")
    return payload[key]


def _text(value: Any, *, where: str, key: str) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise ReviewFormatError(f"{where}: '{key}' metin değil")
    return value


def parse_review(payload: Any) -> ReviewReport:
    """Strictly parse a ``collex.contract-review/v1`` object."""
    if not isinstance(payload, dict):
        raise ReviewFormatError("rapor bir JSON nesnesi değil")
    schema = payload.get("schema", REVIEW_SCHEMA)
    if schema != REVIEW_SCHEMA:
        raise ReviewFormatError(f"desteklenmeyen şema: {schema!r} (beklenen {REVIEW_SCHEMA})")

    findings_raw = _require(payload, "findings", where="rapor")
    if not isinstance(findings_raw, list):
        raise ReviewFormatError("rapor: 'findings' listesi eksik")
    findings: list[ReviewFinding] = []
    for index, entry in enumerate(findings_raw):
        if not isinstance(entry, dict):
            raise ReviewFormatError(f"findings[{index}] bir nesne değil")
        where = f"findings[{index}]"
        numbers_raw = entry.get("clauseNumbers") or []
        if not isinstance(numbers_raw, list):
            raise ReviewFormatError(f"{where}: 'clauseNumbers' liste değil")
        labels_raw = entry.get("clauseLabels") or []
        if not isinstance(labels_raw, list):
            raise ReviewFormatError(f"{where}: 'clauseLabels' liste değil")
        matches_raw = entry.get("matches") or []
        if not isinstance(matches_raw, list):
            raise ReviewFormatError(f"{where}: 'matches' liste değil")
        matches: list[ReviewExcerpt] = []
        for match in matches_raw:
            if not isinstance(match, dict):
                raise ReviewFormatError(f"{where}: 'matches' öğesi nesne değil")
            excerpt = _text(match.get("excerpt"), where=where, key="excerpt")
            if excerpt == "":
                continue
            clause_label = _text(match.get("clauseLabel"), where=where, key="clauseLabel")
            if clause_label == "":
                clause_label = clause_display_label(
                    _text(match.get("clauseNumber"), where=where, key="clauseNumber")
                )
            matches.append(
                ReviewExcerpt(
                    clause_label=clause_label,
                    excerpt=excerpt,
                    negated=match.get("negated") is True,
                    weak=match.get("strength") == "weak",
                )
            )
        findings.append(
            ReviewFinding(
                label=_text(_require(entry, "label", where=where), where=where, key="label"),
                state_label=_text(entry.get("stateLabel"), where=where, key="stateLabel"),
                clause_numbers=tuple(
                    _text(n, where=where, key="clauseNumbers") for n in numbers_raw
                ),
                matched_term=_text(entry.get("matchedTerm"), where=where, key="matchedTerm"),
                note=_text(entry.get("note"), where=where, key="note"),
                clause_labels=tuple(
                    _text(n, where=where, key="clauseLabels") for n in labels_raw
                ),
                reason=_text(entry.get("reason"), where=where, key="reason"),
                matches=tuple(matches),
            )
        )

    clauses_raw = payload.get("clauses") or []
    if not isinstance(clauses_raw, list):
        raise ReviewFormatError("rapor: 'clauses' liste değil")
    observations: list[ReviewObservation] = []
    for index, clause in enumerate(clauses_raw):
        if not isinstance(clause, dict):
            raise ReviewFormatError(f"clauses[{index}] bir nesne değil")
        where = f"clauses[{index}]"
        number = _text(clause.get("clauseNumber"), where=where, key="clauseNumber")
        label = _text(clause.get("label"), where=where, key="label") or clause_display_label(number)
        entries = clause.get("observations") or []
        if not isinstance(entries, list):
            raise ReviewFormatError(f"{where}: 'observations' liste değil")
        for entry in entries:
            if not isinstance(entry, dict):
                raise ReviewFormatError(f"{where}: gözlem nesne değil")
            observations.append(
                ReviewObservation(
                    clause_number=number,
                    text=_text(_require(entry, "text", where=where), where=where, key="text"),
                    sourced=entry.get("sourced") is True,
                    evidence_label=_text(
                        entry.get("evidenceLabel"), where=where, key="evidenceLabel"
                    ),
                    where=label,
                )
            )

    unattached_raw = payload.get("unattachedObservations") or []
    if not isinstance(unattached_raw, list):
        raise ReviewFormatError("rapor: 'unattachedObservations' liste değil")
    unattached: list[ReviewObservation] = []
    for index, entry in enumerate(unattached_raw):
        where = f"unattachedObservations[{index}]"
        if not isinstance(entry, dict):
            raise ReviewFormatError(f"{where} bir nesne değil")
        unattached.append(
            ReviewObservation(
                clause_number="",
                text=_text(_require(entry, "text", where=where), where=where, key="text"),
                sourced=entry.get("sourced") is True,
                evidence_label=_text(entry.get("evidenceLabel"), where=where, key="evidenceLabel"),
                where="Maddeye bağlanamayan gözlem",
                note=_text(entry.get("note"), where=where, key="note"),
            )
        )

    totals_raw = payload.get("totals") or {}
    if not isinstance(totals_raw, dict):
        raise ReviewFormatError("rapor: 'totals' nesne değil")
    totals = {str(k): int(v) for k, v in totals_raw.items() if isinstance(v, int)}

    notices_raw = payload.get("notices") or []
    if not isinstance(notices_raw, list) or not all(isinstance(n, str) for n in notices_raw):
        raise ReviewFormatError("rapor: 'notices' geçersiz")

    clause_count = payload.get("clauseCount", len(clauses_raw))
    if not isinstance(clause_count, int) or clause_count < 0:
        raise ReviewFormatError("rapor: 'clauseCount' geçersiz")

    return ReviewReport(
        document_title=_text(payload.get("documentTitle"), where="rapor", key="documentTitle"),
        checklist_title=_text(payload.get("checklistTitle"), where="rapor", key="checklistTitle"),
        generated_at=_text(payload.get("generatedAt"), where="rapor", key="generatedAt"),
        clause_count=clause_count,
        findings=tuple(findings),
        observations=tuple(observations),
        totals=totals,
        notices=tuple(notices_raw),
        unattached=tuple(unattached),
    )


def load_review(path: str | Path) -> ReviewReport:
    try:
        raw = Path(path).read_text(encoding="utf-8")
    except OSError as exc:
        raise ReviewFormatError(f"rapor dosyası okunamadı: {exc}") from exc
    try:
        payload = json.loads(raw)
    except json.JSONDecodeError as exc:
        raise ReviewFormatError(f"rapor JSON değil: {exc}") from exc
    return parse_review(payload)


# --------------------------------------------------------------------------
# Document construction
# --------------------------------------------------------------------------


def _ensure_style(document: Any, name: str, base: str) -> Any:
    for existing in document.styles:
        if existing.name == name:
            return existing
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles[base]
    return style


def _register(document: Any) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    title = _ensure_style(document, STYLE_TITLE, "Normal")
    title.font.bold = True
    title.font.size = Pt(16)

    warning = _ensure_style(document, STYLE_WARNING, "Normal")
    warning.font.bold = True
    warning.font.color.rgb = RGBColor(0xB0, 0x00, 0x00)

    field = _ensure_style(document, STYLE_FIELD, "Normal")
    field.font.size = Pt(10)
    field.paragraph_format.space_after = Pt(2)

    legend = _ensure_style(document, STYLE_LEGEND, "Normal")
    legend.font.size = Pt(9)
    legend.paragraph_format.space_after = Pt(2)


def _para(document: Any, value: str, style: str | None = None, *, bold: bool = False) -> Any:
    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if value:
        run = paragraph.add_run(value)
        if bold or style in (STYLE_TITLE, STYLE_WARNING):
            run.bold = True
    return paragraph


def build_review_document(report: ReviewReport, *, system_version: str) -> Any:
    document = Document()
    for section in document.sections:
        section.page_width = Cm(A4_WIDTH_CM)
        section.page_height = Cm(A4_HEIGHT_CM)
        section.top_margin = Cm(MARGIN_CM)
        section.bottom_margin = Cm(MARGIN_CM)
        section.left_margin = Cm(MARGIN_CM)
        section.right_margin = Cm(MARGIN_CM)
    _register(document)

    _para(document, REVIEW_BANNER, STYLE_WARNING)
    _para(document, TITLE, STYLE_TITLE)
    if report.document_title:
        _para(document, f"Belge: {report.document_title}", STYLE_FIELD)
    if report.checklist_title:
        _para(document, f"Kontrol listesi: {report.checklist_title}", STYLE_FIELD)
    _para(document, f"İncelenen madde sayısı: {report.clause_count}", STYLE_FIELD)
    if report.generated_at:
        _para(document, f"Rapor tarihi: {T.human_timestamp(report.generated_at)}", STYLE_FIELD)
    _para(
        document,
        "Özet: "
        + " · ".join(
            f"{label}: {report.totals.get(code, 0)}"
            for code, label in (("VAR", "var"), ("YOK", "yok"), ("BELIRSIZ", "belirsiz"))
        ),
        STYLE_FIELD,
        bold=True,
    )

    # ---- checklist table ---------------------------------------------------
    table = document.add_table(rows=1, cols=len(_TABLE_HEADERS))
    table.style = "Table Grid"
    header = table.rows[0].cells
    for index, label in enumerate(_TABLE_HEADERS):
        header[index].text = label
        for paragraph in header[index].paragraphs:
            for run in paragraph.runs:
                run.bold = True

    for finding in report.findings:
        cells = table.add_row().cells
        cells[0].text = finding.label
        cells[1].text = finding.state_label
        # The display label names Özel Şart 1 as "Özel Şartlar 1", never as a
        # bare "1" it shares with Genel Şart 1; older reports carry numbers.
        cells[2].text = ", ".join(
            finding.clause_labels
            if len(finding.clause_labels) == len(finding.clause_numbers)
            else finding.clause_numbers
        )
        cells[3].text = finding.matched_term
        cells[4].text = finding.note

    # ---- excerpts: the contract's own sentence behind every row -----------
    quoted = [finding for finding in report.findings if finding.matches or finding.reason]
    if quoted:
        _para(document, "")
        _para(document, "Metinden alıntılar", STYLE_FIELD, bold=True)
        for finding in quoted:
            if finding.reason:
                _para(document, f"{finding.label} ({finding.state_label}): {finding.reason}", STYLE_FIELD)
            for match in finding.matches:
                marks = []
                if match.negated:
                    marks.append("olumsuz ifade")
                if match.weak:
                    marks.append("şüpheli ifade")
                suffix = f" ({', '.join(marks)})" if marks else ""
                _para(
                    document,
                    f"{finding.label} — {match.clause_label}: “{match.excerpt}”{suffix}",
                    STYLE_LEGEND,
                )

    # ---- observations ------------------------------------------------------
    if report.observations:
        _para(document, "")
        _para(document, "Madde gözlemleri", STYLE_FIELD, bold=True)
        for observation in report.observations:
            where = observation.where or clause_display_label(observation.clause_number)
            if observation.sourced:
                suffix = f" [Kaynak: {observation.evidence_label}]" if observation.evidence_label else ""
                _para(document, f"{where}: {observation.text}{suffix}", STYLE_FIELD)
            else:
                # Already carries the KAYNAKSIZ prefix from the producer; the
                # style makes it visible as well as labelled.
                _para(document, f"{where}: {observation.text}", STYLE_WARNING)

    if report.unattached:
        _para(document, "")
        _para(document, "Maddeye bağlanamayan gözlemler", STYLE_FIELD, bold=True)
        for observation in report.unattached:
            if observation.note:
                _para(document, observation.note, STYLE_LEGEND)
            suffix = (
                f" [Kaynak: {observation.evidence_label}]"
                if observation.sourced and observation.evidence_label
                else ""
            )
            _para(
                document,
                f"{observation.text}{suffix}",
                STYLE_FIELD if observation.sourced else STYLE_WARNING,
            )

    # ---- legend + notices --------------------------------------------------
    _para(document, "")
    _para(document, "Durum sözlüğü", STYLE_FIELD, bold=True)
    for label, meaning in STATE_LEGEND:
        _para(document, f"{label}: {meaning}", STYLE_LEGEND)
    _para(document, UNSOURCED_NOTE, STYLE_LEGEND)
    # The report's notices repeat the banner and the KAYNAKSIZ rule verbatim;
    # a sentence already on the page is not printed a second time.
    printed = {REVIEW_BANNER, UNSOURCED_NOTE}
    for notice in report.notices:
        if notice in printed:
            continue
        printed.add(notice)
        _para(document, notice, STYLE_LEGEND)
    _para(document, f"Üretim: {system_version}", STYLE_LEGEND)
    # Rule 3: the page that closes the report says it is rule-based, too.
    _para(document, REVIEW_BANNER, STYLE_WARNING)
    return document


# --------------------------------------------------------------------------
# Read-back + self-check
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class ReviewDocxReport:
    paragraph_texts: tuple[str, ...]
    table_rows: tuple[tuple[str, ...], ...]
    full_text: str


def read_review_report(path: str | Path) -> ReviewDocxReport:
    document = Document(str(path))
    texts = [p.text for p in document.paragraphs]
    rows: list[tuple[str, ...]] = []
    for table in document.tables:
        for row in table.rows:
            rows.append(tuple(cell.text for cell in row.cells))
    return ReviewDocxReport(
        paragraph_texts=tuple(texts),
        table_rows=tuple(rows),
        full_text="\n".join([*texts, *(" | ".join(r) for r in rows)]),
    )


def _self_check(path: Path, report: ReviewReport) -> None:
    produced = read_review_report(path)
    failures: list[str] = []

    if not produced.paragraph_texts or produced.paragraph_texts[0] != REVIEW_BANNER:
        failures.append("ZORUNLU_UYARI: kural tabanlı bandı ilk paragraf değil")
    if len(produced.table_rows) != len(report.findings) + 1:
        failures.append(
            f"SATIR_SAYISI: tabloda {len(produced.table_rows) - 1} satır var,"
            f" raporda {len(report.findings)}"
        )
    else:
        for finding, written in zip(report.findings, produced.table_rows[1:]):
            if written[0].strip() != finding.label.strip():
                failures.append("BASLIK: kontrol listesi başlığı rapordakiyle birebir değil")
                break
            if written[1].strip() != finding.state_label.strip():
                failures.append("DURUM: durum hücresi rapordakiyle birebir değil")
                break

    for label, _meaning in STATE_LEGEND:
        if label not in produced.full_text:
            failures.append(f"SOZLUK: '{label}' durum sözlüğü belgede yok")
    if UNSOURCED_NOTE not in produced.paragraph_texts:
        failures.append("SOZLUK: KAYNAKSIZ kuralı belgede yazılı değil")

    # RULE 2, on the artifact: no KAYNAKSIZ OBSERVATION may contain the word
    # "risk". The legend and the notices are excluded by identity, not by a
    # pattern: UNSOURCED_NOTE says the word in order to forbid it
    # ("… 'risk' olarak adlandırılmaz"), and a rule that cannot state itself
    # is a rule nobody can read. Everything else with the label is an
    # observation, and an observation may not use the word at all.
    boilerplate = {UNSOURCED_NOTE, *report.notices}
    for line in produced.paragraph_texts:
        if line in boilerplate:
            continue
        if KAYNAKSIZ_PREFIX in line and _RISK_WORD.search(line) is not None:
            failures.append(
                "KAYNAKSIZ_RISK: kaynağa bağlanmamış bir gözlem 'risk' sözcüğünü taşıyor"
            )
            break

    # Every observation the report carries must actually be in the document:
    # a line silently dropped is a review that under-reports itself.
    for observation in (*report.observations, *report.unattached):
        if observation.text not in produced.full_text:
            failures.append("GOZLEM: bir madde gözlemi belgede yok")
            break

    # The excerpt is what lets the lawyer catch a wrong match: one that did
    # not reach the page is a finding the reader cannot check.
    for finding in report.findings:
        if any(match.excerpt not in produced.full_text for match in finding.matches):
            failures.append("ALINTI: bir bulgunun sözleşme alıntısı belgede yok")
            break

    # Each sentence of boilerplate once: a repeated warning is an unread one.
    if produced.paragraph_texts.count(UNSOURCED_NOTE) != 1:
        failures.append("TEKRAR: KAYNAKSIZ kuralı belgede birden çok kez yazılı")
    if produced.paragraph_texts.count(REVIEW_BANNER) != 2:
        failures.append("TEKRAR: kural tabanlı bandı yalnız başta ve sonda olmalı")

    if failures:
        raise ExportRefused("İnceleme raporu kendisiyle doğrulanamadı", failures)


# --------------------------------------------------------------------------
# Public entry point
# --------------------------------------------------------------------------


@dataclass(frozen=True)
class ReviewExportResult:
    path: Path
    format: str
    finding_count: int
    var: int
    yok: int
    belirsiz: int
    unsourced: int

    def summary(self) -> str:
        return " | ".join(
            [
                f"INCELEME-DOCX yazıldı: {self.path}",
                f"başlık: {self.finding_count}",
                f"var: {self.var}",
                f"yok: {self.yok}",
                f"belirsiz: {self.belirsiz}",
                f"kaynaksız gözlem: {self.unsourced}",
            ]
        )


def export_review_docx(
    report: ReviewReport,
    out_path: str | Path,
    *,
    system_version: str | None = None,
) -> ReviewExportResult:
    """Build, write-to-temp, re-open, self-check, then publish."""
    from export import __version__

    version = system_version or f"ColleX export {__version__}"
    document = build_review_document(report, system_version=version)

    target = Path(out_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    handle, tmp_name = tempfile.mkstemp(
        prefix=f".{target.stem}-", suffix=".docx.tmp", dir=str(target.parent)
    )
    os.close(handle)
    tmp_path = Path(tmp_name)
    try:
        document.save(str(tmp_path))
        _self_check(tmp_path, report)
        os.replace(tmp_path, target)
    except BaseException:
        tmp_path.unlink(missing_ok=True)
        raise

    return ReviewExportResult(
        path=target,
        format=FORMAT_NAME,
        finding_count=len(report.findings),
        var=report.totals.get("VAR", 0),
        yok=report.totals.get("YOK", 0),
        belirsiz=report.totals.get("BELIRSIZ", 0),
        unsourced=sum(1 for o in report.observations if not o.sourced),
    )
