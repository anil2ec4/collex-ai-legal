"""W14 phase C · N-3 — the TASLAK copy keeps the same wording discipline.

MEASURED (W14-F-VERIFY §4.4 / §6 N-3, 02.09.2026): the NİHAİ copy carried
zero machine strings, while the TASLAK copy — the one the lawyer reads while
revising — opened a warning with a BARE evidence id:

    Düzenleme uyarısı: 'ev-f4d62195ca825b82' alıntısı paragraf metninde
    birebir bulunamadı: … (QUOTE_ALTERED). …

`ev-f4d62195ca825b82` says nothing to a lawyer, and it was the first thing the
sentence showed. The rule (W13-BACKLOG B-27) is: Turkish first, machine token
in parentheses AT MOST. `(QUOTE_ALTERED)` already obeyed it; the id did not.

All content here is SENTETİK.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document

from export.draft import ExportMode, parse_draft, presentable_warning
from export.petition import export_petition_docx

from tests.export.test_petition_docx import make_draft

STAMP = "2026-09-03T00:00:00+00:00"

#: An evidence id exactly as the control-plane mints it.
EVIDENCE_ID = re.compile(r"\bev-[0-9a-z]{6,64}\b")
#: An UPPER_SNAKE machine code (two segments or more).
MACHINE_CODE = re.compile(r"\b[A-Z][A-Z0-9]*(?:_[A-Z0-9]+)+\b")

#: The exact warning `reviseDraft` writes when the quote-integrity gate fires
#: (control-plane/src/drafting/{revise,quoteIntegrity}.ts).
BROKEN_BINDING_WARNING = (
    "Düzenleme uyarısı: 'ev-f4d62195ca825b82' alıntısı paragraf metninde"
    " birebir bulunamadı: alıntı değiştirildi — kanıt bağı koptu"
    " (QUOTE_ALTERED). Atıf yazılmadı; paragraf KAYNAKSIZ işaretlendi."
    " Alıntıyı metne birebir geri alın veya atfı kaldırın."
)


def _docx_text(path: Path) -> str:
    document = Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def _outside_parentheses(pattern: re.Pattern[str], text: str) -> list[str]:
    """Matches that are NOT inside a parenthetical on their own line."""
    found: list[str] = []
    for match in pattern.finditer(text):
        line_start = text.rfind("\n", 0, match.start()) + 1
        head = text[line_start : match.start()]
        if head.rfind("(") > head.rfind(")"):
            continue  # inside an open parenthesis
        found.append(match.group(0))
    return found


def _export(tmp_path: Path, mode: ExportMode, name: str) -> Path:
    payload = make_draft()
    payload["warnings"].append(BROKEN_BINDING_WARNING)
    target = tmp_path / name
    export_petition_docx(parse_draft(payload), target, generated_at=STAMP, mode=mode)
    return target


class TestPresentableWarning:
    def test_a_quoted_evidence_id_stops_leading_the_sentence(self) -> None:
        out = presentable_warning(BROKEN_BINDING_WARNING)
        # The id is still there — a developer needs it — but in parentheses,
        # after Turkish, and it no longer opens the clause.
        assert "'ev-f4d62195ca825b82'" not in out
        assert "(kanıt kimliği: ev-f4d62195ca825b82)" in out
        assert out.startswith("Düzenleme uyarısı: bir kanıtın (kanıt kimliği:")
        # The rest of the sentence is untouched: this is presentation, not a
        # rewrite of what the gate said.
        assert "alıntı değiştirildi — kanıt bağı koptu (QUOTE_ALTERED)" in out
        assert "paragraf KAYNAKSIZ işaretlendi" in out

    def test_a_known_evidence_id_is_named_by_its_K_number(self) -> None:
        out = presentable_warning(
            "'ev-tck157' alıntısı paragraf metninde birebir bulunamadı.",
            {"ev-tck157": 1},
        )
        assert out.startswith("K-1 numaralı kanıtın (kanıt kimliği: ev-tck157) alıntısı")

    def test_an_already_parenthesised_reference_is_left_alone(self) -> None:
        text = "Kanıt bağı koptu (kanıt kimliği: ev-f4d62195ca825b82)."
        assert presentable_warning(text) == text

    def test_a_bare_id_in_running_text_is_moved_into_parentheses(self) -> None:
        out = presentable_warning("Kanıt ev-f4d62195ca825b82 taslakta yok.")
        assert "(kanıt kimliği: ev-f4d62195ca825b82)" in out
        assert _outside_parentheses(EVIDENCE_ID, out) == []

    def test_a_warning_with_no_evidence_id_is_returned_verbatim(self) -> None:
        text = "1 paragraf KAYNAKSIZ: hukukî dayanağı doğrulanamadı."
        assert presentable_warning(text) == text


class TestMarkedCopyBody:
    def test_the_TASLAK_copy_shows_no_bare_evidence_id(self, tmp_path: Path) -> None:
        text = _docx_text(_export(tmp_path, ExportMode(), "taslak.docx"))
        # The defect: the id led the sentence.
        assert "'ev-f4d62195ca825b82'" not in text
        assert "(kanıt kimliği: ev-f4d62195ca825b82)" in text
        # No machine token stands outside a parenthetical anywhere in the
        # marked copy — the discipline the NİHAİ copy already had.
        assert _outside_parentheses(MACHINE_CODE, text) == []

    def test_the_NIHAI_copy_stays_clean(self, tmp_path: Path) -> None:
        text = _docx_text(
            _export(tmp_path, ExportMode(annex="none", marks="none"), "nihai.docx")
        )
        assert EVIDENCE_ID.search(text) is None
        assert MACHINE_CODE.search(text) is None

    def test_the_normalization_survives_a_round_trip_through_the_CLI(
        self, tmp_path: Path
    ) -> None:
        from export.cli import EXIT_OK, main

        payload = make_draft()
        payload["warnings"].append(BROKEN_BINDING_WARNING)
        source = tmp_path / "taslak.json"
        source.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        out = tmp_path / "cli.docx"
        assert (
            main(
                [
                    "--draft",
                    str(source),
                    "--out",
                    str(out),
                    "--format",
                    "dilekce-docx",
                    "--quiet",
                ]
            )
            == EXIT_OK
        )
        text = _docx_text(out)
        assert "'ev-f4d62195ca825b82'" not in text
        assert "(kanıt kimliği: ev-f4d62195ca825b82)" in text
