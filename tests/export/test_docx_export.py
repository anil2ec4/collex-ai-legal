"""Golden-ish DOCX tests: export, then RE-OPEN the file and interrogate it.

Nothing here trusts the in-memory plan. Every assertion is made against the
bytes that actually landed on disk, read back with ``python-docx``, because
the promise made to a lawyer is about the document they open — not about what
the exporter intended to write.

The body/appendix closure check deliberately parses the document by TEXT
(``[3] ...`` headings and ``Atıflar: ...`` lines) rather than by the custom
styles :mod:`export.bundle_docx` uses internally, so it cannot pass just
because the exporter and its own reader agree with each other.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from docx import Document  # noqa: E402

from export import text as T  # noqa: E402
from export.bundle import parse_bundle  # noqa: E402
from export.bundle_docx import export_docx, read_docx_report  # noqa: E402
from export.errors import ExportRefused  # noqa: E402
from export.plan import build_citation_plan  # noqa: E402

_APPENDIX_HEADING = re.compile(r"\A\[(\d+)\]\s")
_CITATION_LINE = re.compile(r"\A(?:Atıflar|Karşıt kaynaklar):\s*(.*)\Z")
_MARKER = re.compile(r"\[(\d+)\]")


def _independent_parse(path: Path) -> tuple[set[int], list[int]]:
    """Body citation numbers and appendix numbers, parsed from plain text."""
    document = Document(str(path))
    body: set[int] = set()
    appendix: list[int] = []
    for paragraph in document.paragraphs:
        content = paragraph.text.strip()
        heading = _APPENDIX_HEADING.match(content)
        if heading:
            appendix.append(int(heading.group(1)))
            continue
        citation = _CITATION_LINE.match(content)
        if citation:
            body.update(int(m.group(1)) for m in _MARKER.finditer(citation.group(1)))
    return body, appendix


@pytest.fixture()
def exported(qualified_bundle, tmp_path: Path, generated_at: str) -> Path:
    target = tmp_path / "cevap.docx"
    result = export_docx(qualified_bundle, target, generated_at=generated_at)
    assert result.path == target
    assert target.is_file()
    return target


@pytest.fixture()
def exported_abstain(abstain_bundle, tmp_path: Path, generated_at: str) -> Path:
    target = tmp_path / "cekimser.docx"
    export_docx(abstain_bundle, target, generated_at=generated_at)
    return target


def test_body_and_appendix_citations_match_exactly(exported: Path, qualified_bundle):
    """No dangling markers, no orphan appendix entries — both directions."""
    body, appendix = _independent_parse(exported)
    plan = build_citation_plan(qualified_bundle)

    assert body, "gövdede hiç atıf numarası bulunamadı"
    assert set(appendix) == body
    assert set(appendix) == set(plan.numbered_ids)
    assert len(appendix) == len(set(appendix)), "aynı numara iki kez yazılmış"
    assert sorted(appendix) == list(range(1, len(appendix) + 1))


def test_every_appendix_entry_is_reachable_from_a_claim(exported: Path, qualified_bundle):
    """The reverse direction, stated against the bundle's own claim graph."""
    _, appendix = _independent_parse(exported)
    plan = build_citation_plan(qualified_bundle)
    cited_ids = {
        evidence_id
        for claim in qualified_bundle.claims
        for evidence_id in (*claim.evidence_ids, *claim.contrary_evidence_ids)
    }
    assert {plan.numbers[eid] for eid in cited_ids} == set(appendix)


def test_quotes_are_byte_identical_to_the_bundle(exported: Path, qualified_bundle):
    report = read_docx_report(exported)
    plan = build_citation_plan(qualified_bundle)

    assert len(report.quotes) == len(plan.entries)
    for number, entry in plan.entries:
        written = report.quotes[number]
        assert written == entry.quote
        assert written.encode("utf-8") == entry.quote.encode("utf-8")


def test_hashes_version_ids_and_offsets_are_present(exported: Path, qualified_bundle):
    full_text = read_docx_report(exported).full_text
    plan = build_citation_plan(qualified_bundle)

    for _, entry in plan.entries:
        assert entry.quote_sha256 in full_text
        assert entry.content_sha256 in full_text
        assert entry.document_version_id in full_text
        assert f"{entry.locator.start_char}–{entry.locator.end_char}" in full_text


def test_uploaded_document_currentness_prints_the_shared_label(
    qualified_payload, tmp_path: Path, generated_at: str
):
    """W12-B2/API2: NOT_APPLICABLE currentness (an uploaded document) is
    written as the shared label in the re-opened DOCX, never as the raw code."""
    entry = next(e for e in qualified_payload["evidence"] if e["evidenceId"] == "ev-tck-157")
    entry["currentness"]["status"] = "NOT_APPLICABLE"
    entry["currentness"]["score"] = 1.0
    entry["origin"] = "upload"
    bundle = parse_bundle(qualified_payload)
    target = tmp_path / "yukleme.docx"
    export_docx(bundle, target, generated_at=generated_at)

    full_text = read_docx_report(target).full_text
    assert "Güncellik: yüklediğiniz belge — yürürlük değerlendirilemez" in full_text
    assert "NOT_APPLICABLE" not in full_text
    assert "Güncellik: yürürlükte" in full_text


def test_source_metadata_reaches_the_appendix(exported: Path):
    full_text = read_docx_report(exported).full_text
    assert "Yargıtay 15. Ceza Dairesi" in full_text
    assert "E. 2023/4521, K. 2024/1187" in full_text
    # User-facing dates are GG.AA.YYYY (terim sözlüğü); ISO stays in the
    # Teknik künye, never on the human decision-date line.
    assert "Tarih: 12.03.2024" in full_text
    assert "Tarih: 2024-03-12" not in full_text
    assert "Mevzuat No: 5237" in full_text
    assert "Madde: m. 157" in full_text


def test_lawyer_review_footer_is_on_the_document(exported: Path):
    """Brief 11.5: every human-visible page says a lawyer must review this."""
    document = Document(str(exported))
    footers = [
        paragraph.text
        for section in document.sections
        for paragraph in section.footer.paragraphs
    ]
    assert any(T.LAWYER_REVIEW_FOOTER in text for text in footers)
    assert all(
        not section.footer.is_linked_to_previous for section in document.sections
    )
    assert "AVUKAT İNCELEMESİ ZORUNLUDUR" in "\n".join(footers)


def test_synthetic_data_is_labelled_everywhere_visible(exported: Path):
    document = Document(str(exported))
    footers = "\n".join(
        paragraph.text
        for section in document.sections
        for paragraph in section.footer.paragraphs
    )
    body = read_docx_report(exported).full_text

    assert T.SYNTHETIC_FOOTER_TAG in footers
    assert "DENEME VERİSİ" in body
    assert "DENEME VERİSİ — bu kaynak gerçek değildir" in body


def test_confidence_table_has_all_five_dimensions(exported: Path, qualified_bundle):
    document = Document(str(exported))
    headers = [
        [cell.text for cell in table.rows[0].cells] for table in document.tables
    ]
    confidence_header = next(h for h in headers if "Sonuç" in h and "Kaynağın ağırlığı" in h)
    # Shared five-dimension wording across every surface (terim sözlüğü);
    # English keys live only in the Teknik künye.
    for label in (
        "Doğru kaynağa ulaşma",
        "Alıntının sonucu karşılaması",
        "Kaynağın ağırlığı",
        "Metnin güncelliği",
        "Sorunun kapsanması",
    ):
        assert label in confidence_header
    assert "Kaynak bulma (retrieval)" not in confidence_header
    assert "Kapsam (coverage)" not in confidence_header

    confidence_table = next(
        t for t in document.tables if [c.text for c in t.rows[0].cells] == confidence_header
    )
    assert len(confidence_table.rows) == len(qualified_bundle.claims) + 1
    first_claim = qualified_bundle.claims[0]
    values = [c.text for c in confidence_table.rows[1].cells]
    assert values[0] == "1. Sonuç"
    assert values[2] == T.pct(first_claim.confidence.retrieval)


def test_conflicting_authorities_section_names_both_sides(exported: Path):
    report = read_docx_report(exported)
    assert report.has_section(T.H_CONFLICT)
    assert "c2-edimin-ifa-edilmemesi: destekleyen [2] — karşıt [3]" in report.full_text


def test_unsupported_claim_is_flagged_not_hidden(exported: Path):
    full_text = read_docx_report(exported).full_text
    assert "c4-teselsul-iddiasi" in full_text
    assert "KAYNAKSIZ" in full_text
    assert "(doğrulanmış atıf yok)" in full_text


def test_verification_section_explains_re_verification(exported: Path):
    report = read_docx_report(exported)
    assert report.has_section(T.H_VERIFY)
    full_text = report.full_text
    assert "SHA-256" in full_text
    assert "harfi harfine aynı olmalıdır" in full_text
    assert "Kaynağın bu nüshasının numarası" in full_text
    assert "imzalı UDF üretmez" in full_text


def test_abstention_document_has_no_citations_at_all(exported_abstain: Path):
    body, appendix = _independent_parse(exported_abstain)
    report = read_docx_report(exported_abstain)

    assert body == set()
    assert appendix == []
    assert report.quotes == {}
    assert report.has_section(T.H_ABSTAIN)
    assert "numaralandırılmış kaynak YOKTUR (0 atıf)" in report.full_text
    assert T.ABSTENTION_TEXT in report.full_text
    # The considered-but-unpublished sources are counted, not hidden.
    assert "Değerlendirilen ancak yayımlanmayan kaynak sayısı: 1." in report.full_text


def test_abstention_document_still_carries_the_footer(exported_abstain: Path):
    document = Document(str(exported_abstain))
    footers = "\n".join(
        paragraph.text
        for section in document.sections
        for paragraph in section.footer.paragraphs
    )
    assert T.LAWYER_REVIEW_FOOTER in footers


def test_tampered_bundle_never_produces_a_docx(
    qualified_payload, tamper, tmp_path: Path
):
    bundle = parse_bundle(tamper(qualified_payload, 2))
    target = tmp_path / "sahte.docx"
    with pytest.raises(ExportRefused) as excinfo:
        export_docx(bundle, target)

    assert not target.exists()
    assert "QUOTE_HASH_MISMATCH" in excinfo.value.report()


def test_export_is_reproducible_for_a_fixed_timestamp(
    qualified_bundle, tmp_path: Path, generated_at: str
):
    """Same bundle + same stamp -> same visible content (not byte-equal zip)."""
    first = tmp_path / "a.docx"
    second = tmp_path / "b.docx"
    export_docx(qualified_bundle, first, generated_at=generated_at)
    export_docx(qualified_bundle, second, generated_at=generated_at)

    assert read_docx_report(first).full_text == read_docx_report(second).full_text


def test_untrusted_bracket_text_cannot_forge_a_citation(
    qualified_payload, tmp_path: Path, generated_at: str
):
    """A claim whose prose contains ``[9]`` must not create a phantom marker."""
    qualified_payload["claims"][0]["text"] += " (bkz. [9] ve [10])"
    bundle = parse_bundle(qualified_payload)
    target = tmp_path / "enjeksiyon.docx"
    export_docx(bundle, target, generated_at=generated_at)

    body, appendix = _independent_parse(target)
    assert 9 not in body and 10 not in body
    assert set(appendix) == body
