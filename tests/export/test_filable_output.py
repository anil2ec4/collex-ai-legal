"""W14 · B-02 — the filable output, and B-36's review record on the artifact.

W13-DAILYFLOW §3.2 measured the produced dava dilekçesi and could not file it:

  #1  77 % of the Markdown (10 162 of 13 207 B) and 189 of 236 DOCX
      paragraphs were a SHA-256 appendix the body never cited;
  #2  the page was US Letter (7772400 × 10058400 EMU), not A4;
  #3  ``bold=False`` and ``alignment=None`` on 236/236 paragraphs;
  #7  every upload was described as "kiracı yüklemesi";
  #9  the UDF body carried the evidence-bundle schema tag as readable text.

These tests pin the fix. Neither switch moves the gate: the last group proves
``marks=none`` still refuses an unverifiable draft (trap §C.2).

All content is SENTETİK — authored for these tests, not real Turkish law.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pytest

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm

from export.cli import EXIT_OK, EXIT_REFUSED, main
from export.draft import (
    ANNEX_NONE,
    KAYNAKSIZ_PREFIX,
    MARKS_NONE,
    VERIFICATION_INCOMPLETE_LINE,
    ExportMode,
    parse_draft,
)
from export.petition import export_petition_docx, read_petition_report
from export.udf import export_udf, read_udf_report

from tests.export.test_petition_docx import make_contrary_draft, make_draft

FINAL = ExportMode(annex=ANNEX_NONE, marks=MARKS_NONE)
STAMP = "2026-09-02T10:00:00+00:00"

#: The UYAP single-file limit the produced document must stay under.
UYAP_MAX_BYTES = 40 * 1024 * 1024

_SHA256 = re.compile(r"\b[0-9a-f]{64}\b")


def _write(tmp_path: Path, payload: dict) -> Path:
    path = tmp_path / "taslak.json"
    path.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    return path


# --------------------------------------------------------------------------
# Page geometry and typography apply to EVERY mode (DAILYFLOW #2, #3)
# --------------------------------------------------------------------------


@pytest.mark.parametrize("mode", [ExportMode(), FINAL], ids=["taslak", "nihai"])
def test_page_is_a4_with_turkish_margins(tmp_path: Path, mode):
    draft = parse_draft(make_draft())
    out = tmp_path / "dilekce.docx"
    export_petition_docx(draft, out, generated_at=STAMP, mode=mode)

    document = Document(str(out))
    # Word stores page geometry in TWIPS, so a cm value round-trips with a
    # sub-micron rounding difference; compare in cm, not in raw EMU.
    for section in document.sections:
        assert abs(section.page_width.cm - 21.0) < 0.01, "sayfa genişliği A4 değil"
        assert abs(section.page_height.cm - 29.7) < 0.01, "sayfa yüksekliği A4 değil"
        assert abs(section.left_margin.cm - 2.5) < 0.01
        assert abs(section.right_margin.cm - 2.5) < 0.01
        # The measured defect: US Letter (8.5 in = 21.59 cm).
        assert abs(section.page_width.inches - 8.5) > 0.1


@pytest.mark.parametrize("mode", [ExportMode(), FINAL], ids=["taslak", "nihai"])
def test_alignment_and_bold_are_real_on_the_artifact(tmp_path: Path, mode):
    draft = parse_draft(make_draft())
    out = tmp_path / "dilekce.docx"
    export_petition_docx(draft, out, generated_at=STAMP, mode=mode)

    document = Document(str(out))
    alignments = {p.alignment for p in document.paragraphs}
    assert WD_ALIGN_PARAGRAPH.CENTER in alignments, "mahkeme hitabı ortalanmamış"
    assert WD_ALIGN_PARAGRAPH.RIGHT in alignments, "imza bloğu sağa yaslanmamış"
    assert WD_ALIGN_PARAGRAPH.JUSTIFY in alignments, "gövde iki yana yaslanmamış"
    assert any(
        run.bold is True for p in document.paragraphs for run in p.runs
    ), "belgede hiç kalın metin yok"

    # The court address itself is the centred one.
    court = next(p for p in document.paragraphs if p.text.endswith("MAHKEMESİ'NE"))
    assert court.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert court.runs[0].bold is True


# --------------------------------------------------------------------------
# annex=none & marks=none — the clean filing copy (B-02 KABUL)
# --------------------------------------------------------------------------


def test_final_copy_is_clean_body_only(tmp_path: Path):
    draft = parse_draft(make_draft())
    out = tmp_path / "dilekce-nihai.docx"
    result = export_petition_docx(draft, out, generated_at=STAMP, mode=FINAL)
    assert result.mode_label == "NİHAİ"

    produced = read_petition_report(out)
    text = produced.full_text

    for forbidden in (
        KAYNAKSIZ_PREFIX,
        "kiracı yüklemesi",
        "collex.export.",
        "dava-dilekcesi",  # a raw template / rule id
        "Alıntı SHA-256",
        "Kanıt kimliği:",
        "[K-1]",
        "Taslak kimliği",
    ):
        assert forbidden not in text, f"nihai kopyada bulunmamalı: {forbidden!r}"

    assert _SHA256.search(text) is None, "nihai kopyada SHA-256 kaldı"
    assert any(h.endswith("(NİHAİ)") for h in produced.headings)
    assert any(p.startswith("Dayanak: ") for p in produced.paragraph_texts)
    assert out.stat().st_size <= UYAP_MAX_BYTES


def test_final_copy_keeps_every_body_paragraph_and_drops_the_apparatus(tmp_path: Path):
    """The measured defect was 47 body vs 189 annex paragraphs."""
    draft = parse_draft(make_contrary_draft())
    full = tmp_path / "tam.docx"
    final = tmp_path / "nihai.docx"
    export_petition_docx(draft, full, generated_at=STAMP)
    export_petition_docx(draft, final, generated_at=STAMP, mode=FINAL)

    full_report = read_petition_report(full)
    final_report = read_petition_report(final)
    assert len(final_report.paragraph_texts) / len(full_report.paragraph_texts) <= 0.6

    for section in draft.sections:
        if section.section_id == "ek-dogrulama":
            continue
        for paragraph in section.paragraphs:
            first = paragraph.text.split("\n")[0]
            assert first in final_report.paragraph_texts, paragraph.paragraph_id

    # Contrary case law stays loud in EVERY mode — that is contract A.
    assert any("AKSİ yönündedir" in p for p in final_report.paragraph_texts)


def test_full_mode_keeps_every_annex_entry_and_every_mark(tmp_path: Path):
    """Regression: the default mode loses nothing (B-02 KABUL, 2nd half)."""
    draft = parse_draft(make_contrary_draft())
    out = tmp_path / "tam.docx"
    export_petition_docx(draft, out, generated_at=STAMP)
    produced = read_petition_report(out)

    assert produced.numbering == ((1, "ev-tck157"), (2, "ev-karsit-1"))
    assert produced.citation_refs == ("K-1", "K-2")
    assert len(produced.kaynaksiz_texts) == draft.unsupported_count
    assert "DAYANAK KAYNAKLARI" in produced.headings
    for entry in draft.evidence:
        assert f"Alıntının parmak izi: {entry.quote_sha256} (Teknik adı: SHA-256)" in produced.paragraph_texts


def test_upload_verification_line_no_longer_says_kiraci(tmp_path: Path):
    """DAILYFLOW #7 / COPY M1: 'tenant' is a database word, not a party role."""
    payload = make_draft()
    payload["sections"].insert(
        3,
        {
            "id": "deliller",
            "title": "DELİLLER",
            "paragraphs": [
                {
                    "id": "p-deliller-99",
                    "text": (
                        "Ek-1 — kira_sozlesmesi.txt — yüklediğiniz belge, 2 parça"
                        " — SHA-256 (ilk 8): abcdef12"
                        " — Doğrulama: yüklediğiniz belge; hukukî dayanak değildir"
                    ),
                    "evidenceIds": [],
                    "supported": True,
                    "note": "doğrulama bilgisi — belge gövdesine ait değildir",
                    "role": "ekDogrulama",
                }
            ],
        },
    )
    draft = parse_draft(payload)
    out = tmp_path / "dilekce.docx"
    export_petition_docx(draft, out, generated_at=STAMP)
    text = read_petition_report(out).full_text
    assert "yüklediğiniz belge" in text
    assert "kiracı yüklemesi" not in text


# --------------------------------------------------------------------------
# UDF: one format contract across all three outputs
# --------------------------------------------------------------------------


def test_udf_carries_alignment_and_drops_the_schema_tag(tmp_path: Path):
    draft = parse_draft(make_draft())
    out = tmp_path / "dilekce.udf"
    export_udf(draft, out, generated_at=STAMP)
    produced = read_udf_report(out)

    # DAILYFLOW #3: every paragraph used to be Alignment="0".
    assert set(produced.alignments) != {0}, "UDF hizalaması hâlâ tek değer"
    assert 1 in produced.alignments, "mahkeme hitabı ortalanmamış"
    assert 3 in produced.alignments, "gövde iki yana yaslanmamış"
    # DAILYFLOW #9: the evidence-bundle schema tag is not this document's schema.
    assert "collex.export." not in produced.text


def test_udf_final_copy_is_clean(tmp_path: Path):
    draft = parse_draft(make_draft())
    out = tmp_path / "nihai.udf"
    result = export_udf(draft, out, generated_at=STAMP, mode=FINAL)
    assert result.mode_label == "NİHAİ"
    produced = read_udf_report(out)
    assert KAYNAKSIZ_PREFIX not in produced.text
    assert "[K-" not in produced.text
    assert "DAYANAK KAYNAKLARI" not in produced.text
    assert _SHA256.search(produced.text) is None
    assert any(line.startswith("Dayanak: ") for line in produced.lines)


# --------------------------------------------------------------------------
# The gate does not move (trap §C.2)
# --------------------------------------------------------------------------


def test_marks_none_still_refuses_an_unverifiable_draft(tmp_path: Path, capsys):
    """A clean copy of a broken draft is still no copy at all."""
    payload = make_draft()
    payload["evidence"][0]["quoteSha256"] = "0" * 64
    src = _write(tmp_path, payload)
    out = tmp_path / "nihai.docx"

    code = main(
        [
            "--draft", str(src), "--out", str(out),
            "--format", "dilekce-docx",
            "--annex", "none", "--marks", "none",
        ]
    )
    assert code == EXIT_REFUSED
    assert not out.exists()
    assert "ALINTI_OZET_UYUSMAZLIGI" in capsys.readouterr().err


def test_cli_defaults_are_todays_behaviour(tmp_path: Path):
    src = _write(tmp_path, make_draft())
    default_out = tmp_path / "a.docx"
    explicit_out = tmp_path / "b.docx"
    base = ["--draft", str(src), "--format", "dilekce-docx", "--generated-at", STAMP, "--quiet"]
    assert main([*base, "--out", str(default_out)]) == EXIT_OK
    assert main([*base, "--out", str(explicit_out), "--annex", "full", "--marks", "all"]) == EXIT_OK
    assert read_petition_report(default_out) == read_petition_report(explicit_out)


# --------------------------------------------------------------------------
# B-36: the pre-filing review record, on the artifact
# --------------------------------------------------------------------------


def _with_checklist(payload: dict, *, complete: bool) -> dict:
    payload["reviewChecklist"] = {
        "citationsOpened": {
            "checked": True,
            "at": "2026-09-02T09:00:00Z",
            "by": "Av. Ayşe Yılmaz",
        },
        "unsupportedReviewed": {"checked": True, "at": "2026-09-02T09:01:00Z"},
        "contraryRead": {"checked": complete, "at": "2026-09-02T09:02:00Z"},
    }
    return payload


@pytest.mark.parametrize("mode", [ExportMode(), FINAL], ids=["taslak", "nihai"])
def test_unchecked_review_prints_the_incomplete_line(tmp_path: Path, mode):
    draft = parse_draft(make_draft())
    assert draft.review_complete is False
    out = tmp_path / "dilekce.docx"
    export_petition_docx(draft, out, generated_at=STAMP, mode=mode)
    assert VERIFICATION_INCOMPLETE_LINE in read_petition_report(out).paragraph_texts


@pytest.mark.parametrize("mode", [ExportMode(), FINAL], ids=["taslak", "nihai"])
def test_completed_review_removes_the_line(tmp_path: Path, mode):
    draft = parse_draft(_with_checklist(make_draft(), complete=True))
    assert draft.review_complete is True
    out = tmp_path / "dilekce.docx"
    export_petition_docx(draft, out, generated_at=STAMP, mode=mode)
    assert VERIFICATION_INCOMPLETE_LINE not in read_petition_report(out).paragraph_texts


def test_partial_review_still_prints_the_line_in_udf(tmp_path: Path):
    draft = parse_draft(_with_checklist(make_draft(), complete=False))
    assert draft.review_complete is False
    out = tmp_path / "dilekce.udf"
    export_udf(draft, out, generated_at=STAMP)
    assert VERIFICATION_INCOMPLETE_LINE in read_udf_report(out).lines


def test_review_record_survives_the_parser_with_its_note():
    payload = _with_checklist(make_draft(), complete=True)
    payload["evidenceReview"] = {
        "ev-tck157": {
            "checked": True,
            "at": "2026-09-02T09:05:00Z",
            "by": "Av. Ayşe Yılmaz",
            "note": "Madde metni resmî kaynaktan teyit edildi.",
        }
    }
    draft = parse_draft(payload)
    assert draft.review_checklist["citationsOpened"].by == "Av. Ayşe Yılmaz"
    assert draft.evidence_review["ev-tck157"].note.startswith("Madde metni")
    # Unknown checklist keys are IGNORED, never fatal (additive contract).
    payload["reviewChecklist"]["gelecektekiKutu"] = {"checked": True}
    assert parse_draft(payload).review_complete is True


def test_old_draft_without_the_record_still_parses_and_exports(tmp_path: Path):
    payload = make_draft()
    assert "reviewChecklist" not in payload
    draft = parse_draft(payload)
    assert draft.review_checklist == {}
    assert draft.evidence_review == {}
    out = tmp_path / "dilekce.docx"
    assert export_petition_docx(draft, out, generated_at=STAMP).path == out
