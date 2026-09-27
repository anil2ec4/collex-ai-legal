"""Test helper: edit an exported ColleX DOCX the way a lawyer does in Word.

Used by ``tests/export/test_draft_docx_roundtrip.py`` (imported) and by
``control-plane/tests/drafting/docxImport.real.test.ts`` (run as a script):

    python tests/export/docx_word_edit.py <in.docx> <out.docx> <ops.json>

Every edit goes through python-docx / raw WordprocessingML and the file is
saved by python-docx — the same "re-save by another program" a Word save is.
Operations (a JSON list, applied in order; ``contains`` picks the FIRST body
paragraph whose accepted text contains the string):

* ``{"op": "replace_text", "contains": s, "new": t}`` — the paragraph's whole
  text becomes ``t`` (runs rewritten; hidden bookmarks untouched).
* ``{"op": "replace_in", "contains": s, "old": a, "new": b}`` — ``a`` → ``b``
  once inside the paragraph (e.g. one letter of a quote).
* ``{"op": "insert_after", "contains": s, "text": t}`` — a new paragraph in
  the same style right after it (Word's Enter at the end of a paragraph).
* ``{"op": "delete", "contains": s}`` — the paragraph is removed.
* ``{"op": "tracked_insert", "contains": s, "text": t}`` — a PENDING tracked
  insertion (``w:ins``) appended to the paragraph.
* ``{"op": "tracked_delete", "contains": s, "word": w}`` — a PENDING tracked
  deletion (``w:del``) of ``w`` inside the paragraph.
* ``{"op": "resave"}`` — nothing; the file is just opened and saved.
"""

from __future__ import annotations

import copy
import json
import sys
from pathlib import Path
from typing import Any

_REPO_ROOT = Path(__file__).resolve().parents[2]
if str(_REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(_REPO_ROOT))

from docx import Document  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from export.draft_identity import _local_accepted_text  # noqa: E402

_REV_ATTRS = {"w:id": "901", "w:author": "Av. Deneme", "w:date": "2026-09-27T10:00:00Z"}


def _find(document: Any, contains: str) -> Any:
    for paragraph in document.paragraphs:
        if contains in _local_accepted_text(paragraph._p):
            return paragraph
    raise LookupError(f"paragraf bulunamadı: {contains!r}")


def _rewrite(paragraph: Any, text: str) -> None:
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._r.getparent().remove(run._r)


def _rev(tag: str, n: int) -> Any:
    element = OxmlElement(tag)
    for key, value in _REV_ATTRS.items():
        element.set(qn(key), value if key != "w:id" else str(900 + n))
    return element


def apply_ops(document: Any, ops: list[dict[str, Any]]) -> None:
    for n, op in enumerate(ops):
        kind = op["op"]
        if kind == "resave":
            continue
        paragraph = _find(document, op["contains"])
        if kind == "replace_text":
            _rewrite(paragraph, op["new"])
        elif kind == "replace_in":
            text = "".join(run.text for run in paragraph.runs)
            if op["old"] not in text:
                raise LookupError(f"metin yok: {op['old']!r}")
            _rewrite(paragraph, text.replace(op["old"], op["new"], 1))
        elif kind == "insert_after":
            new_p = OxmlElement("w:p")
            ppr = paragraph._p.pPr
            if ppr is not None:
                new_p.append(copy.deepcopy(ppr))
            paragraph._p.addnext(new_p)
            run = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = op["text"]
            t.set(qn("xml:space"), "preserve")
            run.append(t)
            new_p.append(run)
        elif kind == "delete":
            paragraph._p.getparent().remove(paragraph._p)
        elif kind == "tracked_insert":
            ins = _rev("w:ins", n)
            run = OxmlElement("w:r")
            t = OxmlElement("w:t")
            t.text = op["text"]
            t.set(qn("xml:space"), "preserve")
            run.append(t)
            ins.append(run)
            paragraph._p.append(ins)
        elif kind == "tracked_delete":
            word = op["word"]
            runs = paragraph.runs
            text = "".join(run.text for run in runs)
            at = text.find(word)
            if at < 0:
                raise LookupError(f"kelime yok: {word!r}")
            _rewrite(paragraph, text[:at])
            first = paragraph.runs[0]._r
            deleted = _rev("w:del", n)
            del_run = OxmlElement("w:r")
            del_text = OxmlElement("w:delText")
            del_text.text = word
            del_text.set(qn("xml:space"), "preserve")
            del_run.append(del_text)
            deleted.append(del_run)
            first.addnext(deleted)
            tail = OxmlElement("w:r")
            tail_t = OxmlElement("w:t")
            tail_t.text = text[at + len(word):]
            tail_t.set(qn("xml:space"), "preserve")
            tail.append(tail_t)
            deleted.addnext(tail)
        else:
            raise ValueError(f"bilinmeyen işlem: {kind}")


def edit_docx(source: str | Path, target: str | Path, ops: list[dict[str, Any]]) -> None:
    document = Document(str(source))
    apply_ops(document, ops)
    document.save(str(target))


def main(argv: list[str]) -> int:
    source, target, ops_path = argv
    edit_docx(source, target, json.loads(Path(ops_path).read_text(encoding="utf-8")))
    return 0


if __name__ == "__main__":  # pragma: no cover
    sys.exit(main(sys.argv[1:]))
