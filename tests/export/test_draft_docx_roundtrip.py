"""W22 "Word'de düzelttim, geri yükle": the hidden draft identity + read-back.

The export side: every DOCX draft export (all four modes) carries a
tamper-evident ``collex.draft-identity/v1`` custom XML part and hidden
``_cxp<n>`` bookmarks, WITHOUT a single visible character changing. The read
side: ``export.draft_identity.read_draft_docx`` finds the identity wherever
the saving program put it, verifies its digest, and returns the body as it
reads with every tracked change accepted — plus the count of changes still
pending. The alignment/diff and the gates are the control-plane's (vitest:
``control-plane/tests/drafting/docxImport*.test.ts``).
"""

from __future__ import annotations

import hashlib
import json
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from docx_word_edit import edit_docx
from export.draft import ExportMode, canonical_quote_text, parse_draft
from export.draft_identity import (
    BOOKMARK_PREFIX,
    IDENTITY_NS,
    STATUS_DIGEST_MISMATCH,
    STATUS_MISSING,
    STATUS_OK,
    read_draft_docx,
)
from export.petition import export_petition_docx, read_petition_report
from test_petition_docx import QUOTE, make_contrary_draft

_REPO_ROOT = Path(__file__).resolve().parents[2]
MODES = [("full", "all"), ("none", "all"), ("full", "none"), ("none", "none")]


def _payload(annex: str = "full", marks: str = "all") -> dict:
    payload = make_contrary_draft()
    if annex == "none" and marks == "all":
        # Pre-existing rule (B-02): an annex-free copy may not carry a screen
        # mark, so this mode is only producible from a draft with no
        # KAYNAKSIZ paragraph — drop the fixture's one unsourced ground.
        payload["sections"] = [s for s in payload["sections"] if s["id"] != "hukuki-sebepler"]
        payload["unsupportedCount"] = 0
    return payload


def _export(tmp_path: Path, annex: str = "full", marks: str = "all", name: str = "taslak.docx") -> Path:
    out = tmp_path / name
    export_petition_docx(parse_draft(_payload(annex, marks)), out, mode=ExportMode(annex=annex, marks=marks))
    return out


def _identity_item(archive: zipfile.ZipFile) -> str:
    """The customXml item carrying OUR namespace (python-docx's template
    already ships a bibliography item1.xml, so ours is not item1)."""
    hits = [
        name
        for name in archive.namelist()
        if re.fullmatch(r"customXml/item\d+\.xml", name)
        and IDENTITY_NS in archive.read(name).decode("utf-8")
    ]
    assert len(hits) == 1, hits
    return hits[0]


def _body_ids(annex: str, marks: str = "all") -> list[str]:
    payload = _payload(annex, marks)
    return [
        p["id"]
        for s in payload["sections"]
        if not (annex == "none" and s["id"] == "ek-dogrulama")
        for p in s["paragraphs"]
    ]


@pytest.mark.parametrize("annex,marks", MODES)
def test_identity_is_embedded_in_every_mode(tmp_path: Path, annex: str, marks: str):
    out = _export(tmp_path, annex, marks)
    readback = read_draft_docx(out)
    assert readback["identityStatus"] == STATUS_OK
    identity = readback["identity"]
    assert identity["draftId"] == "dft-test-0001"
    assert identity["version"] == 1
    assert identity["exportMode"] == {"annex": annex, "marks": marks}
    assert [p["id"] for p in identity["paragraphs"]] == _body_ids(annex, marks)
    by_id = {p["id"]: p for p in identity["paragraphs"]}
    quote_sha = hashlib.sha256(QUOTE.encode("utf-8")).hexdigest()
    assert by_id["p-aciklamalar-5"]["quotes"] == [{"evidenceId": "ev-tck157", "quoteSha256": quote_sha}]
    if "p-sebepler-6" in by_id:
        assert by_id["p-sebepler-6"]["supported"] is False
    # The paragraph fingerprint is the canonical-text sha256 both runtimes compute.
    text = make_contrary_draft()["sections"][2]["paragraphs"][1]["text"]
    assert by_id["p-aciklamalar-5"]["textSha256"] == hashlib.sha256(
        canonical_quote_text(text).encode("utf-8")
    ).hexdigest()
    # Every paragraph's text lines are bookmarked exactly once.
    marked = [b for b in readback["blocks"] if b["marks"]]
    assert sorted({n for b in marked for n in b["marks"]}) == list(range(1, len(identity["paragraphs"]) + 1))
    # The OOXML data store Word keeps: item + itemProps, related from the document.
    with zipfile.ZipFile(out) as archive:
        item = _identity_item(archive)
        number = re.search(r"(\d+)\.xml$", item).group(1)
        assert f"customXml/itemProps{number}.xml" in archive.namelist()
        rels = archive.read("word/_rels/document.xml.rels").decode("utf-8")
        assert "relationships/customXml" in rels and item in rels
        props = archive.read(f"customXml/_rels/item{number}.xml.rels").decode("utf-8")
        assert "relationships/customXmlProps" in props


@pytest.mark.parametrize("annex,marks", MODES)
def test_identity_adds_no_visible_text(tmp_path: Path, annex: str, marks: str):
    out = _export(tmp_path, annex, marks)
    report = read_petition_report(out)
    visible = "\n".join((report.full_text, *report.footer_texts))
    for needle in ("draft-identity", BOOKMARK_PREFIX):
        assert needle not in visible
    if annex == "none":
        # The filing copy never printed the draft id; the identity must not either.
        assert "dft-test-0001" not in visible
    with zipfile.ZipFile(out) as archive:
        body = archive.read("word/document.xml").decode("utf-8")
    # Hidden bookmarks only: Word hides names that start with "_".
    names = re.findall(r'w:bookmarkStart[^>]*w:name="([^"]+)"', body)
    assert names and all(name.startswith("_cxp") and len(name) <= 40 for name in names)
    # No text run carries anything from the identity.
    runs = "".join(re.findall(r"<w:t(?: [^>]*)?>([^<]*)</w:t>", body))
    assert "draft-identity" not in runs and "_cxp" not in runs


def test_identity_survives_a_resave_by_another_program(tmp_path: Path):
    out = _export(tmp_path)
    before = read_draft_docx(out)
    resaved = tmp_path / "yeniden.docx"
    edit_docx(out, resaved, [{"op": "resave"}])
    after = read_draft_docx(resaved)
    assert after["identityStatus"] == STATUS_OK
    assert after["identity"] == before["identity"]
    assert [b["marks"] for b in after["blocks"]] == [b["marks"] for b in before["blocks"]]


def _rewrite_item(src: Path, dst: Path, transform) -> None:
    with zipfile.ZipFile(src) as source, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as target:
        item = _identity_item(source)
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == item:
                data = transform(data.decode("utf-8")).encode("utf-8")
            target.writestr(info, data)


def test_an_edited_identity_is_detected(tmp_path: Path):
    out = _export(tmp_path)
    tampered = tmp_path / "kurcalanmis.docx"
    _rewrite_item(out, tampered, lambda xml: xml.replace('version="1"', 'version="2"', 1))
    readback = read_draft_docx(tampered)
    assert readback["identityStatus"] == STATUS_DIGEST_MISMATCH


def test_a_docx_without_identity_says_so(tmp_path: Path):
    plain = tmp_path / "duz.docx"
    document = Document()
    document.add_paragraph("Avukatın kendi yazdığı bir dilekçe.")
    document.save(str(plain))
    readback = read_draft_docx(plain)
    assert readback["identityStatus"] == STATUS_MISSING
    assert readback["identity"] is None
    assert readback["blocks"][0]["text"] == "Avukatın kendi yazdığı bir dilekçe."


def test_readback_follows_word_edits_and_reads_tracked_changes_as_accepted(tmp_path: Path):
    out = _export(tmp_path)
    edited = tmp_path / "duzeltilmis.docx"
    edit_docx(
        out,
        edited,
        [
            {"op": "replace_text", "contains": "yatırım vaadinde", "new": "Davalı, davacıya yazılı bir yatırım vaadinde bulunmuştur."},
            {"op": "replace_in", "contains": "Doğrulanmış kaynak uyarınca", "old": "beş yıla", "new": "beş yıle"},
            {"op": "insert_after", "contains": "Faiz başlangıcına", "text": "Faiz, temerrüt tarihinden işler."},
            {"op": "delete", "contains": "DAVALI : Veli Kaya"},
            {"op": "tracked_insert", "contains": "Sentetik alacağın", "text": " (yasal faiziyle)"},
            {"op": "tracked_delete", "contains": "Sentetik alacağın", "word": "davalıdan "},
        ],
    )
    readback = read_draft_docx(edited)
    assert readback["identityStatus"] == STATUS_OK
    blocks = readback["blocks"]
    by_text = {b["text"]: b for b in blocks}
    ids = {p["n"]: p["id"] for p in readback["identity"]["paragraphs"]}

    changed = by_text["Davalı, davacıya yazılı bir yatırım vaadinde bulunmuştur."]
    assert [ids[n] for n in changed["marks"]] == ["p-aciklamalar-4"]  # still found by its bookmark
    altered = next(b for b in blocks if "beş yıle" in b["text"])
    assert [ids[n] for n in altered["marks"]] == ["p-aciklamalar-5"]
    added = by_text["Faiz, temerrüt tarihinden işler."]
    assert added["marks"] == []  # a new paragraph has no identity
    assert all("p-taraflar-3" not in [ids[n] for n in b["marks"]] for b in blocks)  # deleted
    tracked = next(b for b in blocks if "Sentetik alacağın" in b["text"])
    # Accepted view: the insertion is in, the deletion is out.
    assert tracked["text"] == "1. Sentetik alacağın tahsiline karar verilmesini talep ederiz. (yasal faiziyle)"
    assert tracked["pendingChange"] is True
    assert readback["trackedChanges"]["pending"] is True
    assert readback["trackedChanges"]["insertions"] >= 1
    assert readback["trackedChanges"]["deletions"] >= 1


def test_a_tracked_deletion_of_a_paragraph_mark_joins_the_two_paragraphs(tmp_path: Path):
    out = _export(tmp_path)
    document = Document(str(out))
    target = next(p for p in document.paragraphs if p.text == "DAVACI : Ayşe Yılmaz")
    ppr = target._p.get_or_add_pPr()
    rpr = OxmlElement("w:rPr")
    deleted = OxmlElement("w:del")
    deleted.set(qn("w:id"), "950")
    deleted.set(qn("w:author"), "Av. Deneme")
    rpr.append(deleted)
    ppr.append(rpr)
    edited = tmp_path / "birlesik.docx"
    document.save(str(edited))
    readback = read_draft_docx(edited)
    merged = next(b for b in readback["blocks"] if b["text"].startswith("DAVACI : Ayşe Yılmaz"))
    assert merged["text"] == "DAVACI : Ayşe YılmazDAVALI : Veli Kaya"
    assert len(merged["marks"]) == 2 and merged["pendingChange"] is True


def _cli(*args: str) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [sys.executable, "-X", "utf8", "-m", "export.draft_identity", *args],
        cwd=_REPO_ROOT,
        capture_output=True,
        text=True,
        encoding="utf-8",
        timeout=120,
    )


def test_cli_reads_back_after_the_upload_quarantine(tmp_path: Path):
    out = _export(tmp_path)
    done = _cli("--read", str(out), "--name", "Dava Dilekçesi - v1.docx")
    assert done.returncode == 0, done.stderr
    payload = json.loads(done.stdout.strip().splitlines()[-1])
    assert payload["identityStatus"] == STATUS_OK
    assert payload["fileSha256"] == hashlib.sha256(out.read_bytes()).hexdigest()


def test_cli_refuses_what_is_not_a_word_document(tmp_path: Path):
    text = tmp_path / "not.txt"
    text.write_text("düz metin", encoding="utf-8")
    done = _cli("--read", str(text), "--name", "not.txt")
    assert done.returncode == 1
    error = json.loads(done.stdout.strip().splitlines()[-1])["error"]
    assert error["kind"] == "UNSUPPORTED_TYPE"
    assert "Word belgesi" in error["message"]

    disguised = tmp_path / "sahte.docx"
    disguised.write_bytes(b"%PDF-1.4\n%sahte\n")
    done = _cli("--read", str(disguised), "--name", "sahte.docx")
    assert done.returncode == 1
    assert json.loads(done.stdout.strip().splitlines()[-1])["error"]["kind"] == "UNSUPPORTED_TYPE"
