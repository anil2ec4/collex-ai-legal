"""Generator for the SYNTHETIC upload fixtures in this directory.

Run (repo venv only):

    .venv/Scripts/python.exe evals/fixtures/uploads/_generate.py

Every file this script writes is SENTETİK (authored test data, no real
person, case or matter) and says so inside its own content. The files:

* ``dilekce_ornek.docx``  — realistic dava dilekçesi (python-docx):
  taraflar/olaylar/talep, TCK m. 157 + "5237 sayılı" references, an E./K.
  pair, dates in both numeric and Turkish long form, a delil table.
* ``sozlesme_ornek.pdf``  — born-digital single-page PDF authored BY HAND
  as raw bytes (no PDF-writer dependency): Helvetica + WinAnsiEncoding
  with a /Differences array mapping ş ğ ı İ Ş Ğ so Turkish text survives
  pypdf extraction exactly (verified by tests/intake).
* ``metin_ornek.txt``     — UTF-8 plain text with icra content (İİK 89).
* ``udf_ornek.udf``       — minimal synthetic UDF: a zip whose
  ``content.xml`` carries the text in CDATA (UYAP's documented
  zip+XML shape; real signed UDFs are handled read-only, brief 11.3).
* ``bozuk.pdf``           — %PDF magic followed by garbage (parse failure).
* ``taranmis.pdf``        — valid one-page PDF with NO text operators at
  all (the born-digital text layer is empty, like a scan): trips the
  fail-closed OCR gate.
* ``zip_bomba.docx``      — a DOCX-shaped zip whose word/document.xml
  declares 60 MB of zeros (~100x+ compression): trips the per-entry
  uncompressed cap and the ratio cap in intake/quarantine.py.

Zip entries use a FIXED timestamp so regeneration is deterministic.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

HERE = Path(__file__).resolve().parent

ZIP_DATE = (2026, 1, 1, 0, 0, 0)

SENTETIK = "SENTETİK TEST BELGESİDİR — gerçek kişi, olay veya dosya içermez."


# ---------------------------------------------------------------------------
# Hand-authored PDF (WinAnsi + /Differences for Turkish glyphs)
# ---------------------------------------------------------------------------

_TURKISH_DIFF = {
    "ş": 128, "ğ": 129, "ı": 130, "İ": 131, "Ş": 132, "Ğ": 133,
}
_GLYPH_NAMES = ["scedilla", "gbreve", "dotlessi", "Idotaccent",
                "Scedilla", "Gbreve"]


def _encode_pdf_text(line: str) -> bytes:
    out = bytearray()
    for ch in line:
        if ch in _TURKISH_DIFF:
            out.append(_TURKISH_DIFF[ch])
            continue
        try:
            encoded = ch.encode("cp1252")
        except UnicodeEncodeError:
            encoded = b"?"
        for byte in encoded:
            if byte in (0x28, 0x29, 0x5C):  # ( ) \ need escaping
                out.append(0x5C)
            out.append(byte)
    return bytes(out)


def build_pdf(lines: list[str]) -> bytes:
    """Minimal valid single-page PDF; empty ``lines`` -> no text layer."""
    stream = bytearray()
    if lines:
        stream += b"BT\n/F1 11 Tf\n14 TL\n50 780 Td\n"
        for line in lines:
            stream += b"(" + _encode_pdf_text(line) + b") Tj\nT*\n"
        stream += b"ET\n"

    diff = " ".join(["128"] + ["/" + g for g in _GLYPH_NAMES])
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842]"
         b" /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"),
        (b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n"
         + bytes(stream) + b"endstream"),
        (b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica"
         b" /Encoding 6 0 R >>"),
        (b"<< /Type /Encoding /BaseEncoding /WinAnsiEncoding /Differences ["
         + diff.encode() + b"] >>"),
    ]

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for off in offsets:
        out += f"{off:010d} 00000 n \n".encode()
    out += (f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_pos}\n%%EOF\n").encode()
    return bytes(out)


SOZLESME_LINES = [
    "SENTETİK HİZMET SÖZLEŞMESİ",
    SENTETIK,
    "",
    "İşbu sözleşme 6098 sayılı Türk Borçlar Kanunu hükümleri uyarınca",
    "01.02.2025 tarihinde aşağıdaki taraflar arasında akdedilmiştir.",
    "",
    "İŞ SAHİBİ : Ayşe Yılmaz",
    "YÜKLENİCİ : Mehmet Demir",
    "",
    "Madde 1 - Yüklenici, hizmet bedelinin ödenmemesi halinde",
    "TBK 117 uyarınca temerrüt faizi talep edebilir.",
    "Madde 2 - Uyuşmazlık halinde Ankara mahkemeleri yetkilidir.",
]


# ---------------------------------------------------------------------------
# DOCX dilekçe (python-docx)
# ---------------------------------------------------------------------------

def build_dilekce_docx() -> bytes:
    import docx

    document = docx.Document()
    p = document.add_paragraph
    p(SENTETIK)
    p("ANKARA NÖBETÇİ ASLİYE HUKUK MAHKEMESİNE")
    p("")
    p("DAVACI : Ayşe Yılmaz")
    p("VEKİLİ : Av. Mehmet Demir")
    p("DAVALI : Mustafa Kaya")
    p("KONU : Maddi ve manevi tazminat istemine ilişkindir.")
    p("")
    p("AÇIKLAMALAR")
    p("1. Davalı, 12.05.2024 tarihinde davacıya ait iş yerinde "
      "kendisini kamu görevlisi olarak tanıtarak davacıyı aldatmış ve "
      "45.000 TL tutarında haksız menfaat temin etmiştir.")
    p("2. Davalının bu eylemi 5237 sayılı Türk Ceza Kanunu kapsamında "
      "değerlendirilmiş; TCK m. 157 uyarınca dolandırıcılık suçundan "
      "Ankara Cumhuriyet Başsavcılığına şikayette bulunulmuştur.")
    p("3. Yargıtay 9. Hukuk Dairesi'nin E. 2021/123, K. 2022/456 sayılı "
      "kararında benzer nitelikteki aldatma eylemlerinin haksız fiil "
      "sorumluluğu doğurduğu kabul edilmiştir.")
    p("4. Davacı, 3 Haziran 2024 tarihinde davalıya ihtarname göndermiş, "
      "ancak ödeme yapılmamıştır.")
    p("")
    p("DELİLLER")
    table = document.add_table(rows=4, cols=3)
    header = table.rows[0].cells
    header[0].text = "No"
    header[1].text = "Delil"
    header[2].text = "Tarih"
    rows = [
        ("1", "Banka dekontu (sentetik)", "12.05.2024"),
        ("2", "İhtarname (sentetik)", "03.06.2024"),
        ("3", "Tanık beyanı (sentetik)", "-"),
    ]
    for i, (no, delil, tarih) in enumerate(rows, start=1):
        cells = table.rows[i].cells
        cells[0].text = no
        cells[1].text = delil
        cells[2].text = tarih
    p("")
    p("SONUÇ VE İSTEM : Yukarıda açıklanan nedenlerle davanın kabulü ile "
      "45.000 TL maddi ve 10.000 TL manevi tazminatın dava tarihinden "
      "itibaren işleyecek yasal faizi ile birlikte davalıdan tahsiline "
      "karar verilmesini saygıyla vekaleten talep ederiz.")
    p("")
    p("Davacı Vekili")
    p("Av. Mehmet Demir")

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# TXT + UDF + adversarial fixtures
# ---------------------------------------------------------------------------

METIN_TXT = f"""{SENTETIK}

İCRA TAKİBİNE İTİRAZIN İPTALİ TALEBİDİR

ALACAKLI : Fatma Şahin
BORÇLU : Ali Vural

Borçlu aleyhine 15/03/2024 tarihinde başlatılan ilamsız icra takibine
borçlu haksız olarak itiraz etmiştir. İİK 89 kapsamında üçüncü
kişilerdeki hak ve alacaklara haciz ihbarnamesi gönderilmesi de
istenmiştir.

2004 sayılı İcra ve İflas Kanunu hükümleri uyarınca itirazın iptali ile
takibin devamına ve alacağın tahsiline karar verilmesini talep ederiz.
"""

UDF_CONTENT_TEXT = f"""{SENTETIK}

SENTETİK CEVAP DİLEKÇESİ

DAVALI : Mustafa Kaya
VEKİLİ : Av. Zeynep Aksoy

Davacı tarafın 12.01.2025 tarihli dava dilekçesindeki iddialar yerinde
değildir. Müvekkil hakkında TCK m. 157 kapsamında yürütülen soruşturma
delil yetersizliğinden takipsizlikle sonuçlanmıştır.

Açıklanan nedenlerle haksız davanın reddine karar verilmesini talep
ederiz.
"""


def build_udf(text: str) -> bytes:
    content_xml = (
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<template format_id="1.8">\n'
        "<content><![CDATA[" + text + "]]></content>\n"
        '<properties pageFormat="A4" />\n'
        "</template>\n"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        info = zipfile.ZipInfo("content.xml", date_time=ZIP_DATE)
        # writestr with a ZipInfo takes the compression from the INFO
        # (default ZIP_STORED), not from the archive constructor.
        info.compress_type = zipfile.ZIP_DEFLATED
        zf.writestr(info, content_xml.encode("utf-8"))
    return buf.getvalue()


def build_zip_bomba_docx() -> bytes:
    """DOCX-shaped archive declaring 60 MB of zeros in word/document.xml."""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as zf:
        ct = zipfile.ZipInfo("[Content_Types].xml", date_time=ZIP_DATE)
        ct.compress_type = zipfile.ZIP_DEFLATED
        zf.writestr(ct, b"<Types/>")
        doc = zipfile.ZipInfo("word/document.xml", date_time=ZIP_DATE)
        doc.compress_type = zipfile.ZIP_DEFLATED
        zf.writestr(doc, b"\x00" * (60 * 1024 * 1024))
    return buf.getvalue()


def main() -> None:
    outputs = {
        "dilekce_ornek.docx": build_dilekce_docx(),
        "sozlesme_ornek.pdf": build_pdf(SOZLESME_LINES),
        "metin_ornek.txt": METIN_TXT.encode("utf-8"),
        "udf_ornek.udf": build_udf(UDF_CONTENT_TEXT),
        "bozuk.pdf": b"%PDF-1.4\n" + bytes(range(256)) * 16,
        "taranmis.pdf": build_pdf([]),
        "zip_bomba.docx": build_zip_bomba_docx(),
    }
    for name, data in outputs.items():
        target = HERE / name
        # write_bytes: exact bytes, no newline translation on Windows.
        target.write_bytes(data)
        print(f"wrote {name}  ({len(data)} bytes)")


if __name__ == "__main__":
    main()
